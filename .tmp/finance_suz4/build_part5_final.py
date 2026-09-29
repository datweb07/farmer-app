from pathlib import Path
from math import ceil

from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"D:\final app\final app")
XLSX = ROOT / "docs" / "BÁO CÁO TÀI CHÍNH SUZ2026  (4).xlsx"
OUT = ROOT / "docs" / "Phần 5 final.docx"
TMP = ROOT / ".tmp" / "finance_suz4" / "charts"
TMP.mkdir(parents=True, exist_ok=True)

NAVY = "163A5F"
BLUE = "1976A3"
GREEN = "24845B"
ORANGE = "D9822B"
RED = "C73E3A"
PURPLE = "7251A3"
GRAY = "6B7280"
LIGHT = "EDF4F8"
TEXT = "1F2937"
FONT = "Aptos"
FONT_REG = r"C:\Windows\Fonts\arial.ttf"
FONT_BOLD = r"C:\Windows\Fonts\arialbd.ttf"


def img_font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def fmt_m(v, digits=0):
    if v is None or v == "":
        return "-"
    if isinstance(v, str):
        return v
    s = f"{abs(v):,.{digits}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"({s})" if v < 0 else s


def fmt_pct(v, digits=1):
    return f"{v * 100:.{digits}f}%".replace(".", ",")


def set_font(run, size=10.5, bold=False, color=TEXT, italic=False):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    for key in ("ascii", "hAnsi", "eastAsia"):
        rpr.rFonts.set(qn(f"w:{key}"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    node = tcpr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tcpr.append(node)
    node.set(qn("w:fill"), fill)


def borders(cell, color="D7DEE5"):
    tcpr = cell._tc.get_or_add_tcPr()
    box = tcpr.first_child_found_in("w:tcBorders")
    if box is None:
        box = OxmlElement("w:tcBorders")
        tcpr.append(box)
    for edge in ("top", "left", "bottom", "right"):
        el = box.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            box.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)


def cell_margin(cell, val=90):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for side in ("top", "start", "bottom", "end"):
        el = mar.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            mar.append(el)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trpr.append(el)


def add_p(doc, text, bold_lead=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.14
    p.paragraph_format.space_after = Pt(after)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, size=10.2)
    return p


def add_table(doc, headers, rows, widths=None, size=8.7):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    repeat_header(t.rows[0])
    no_split(t.rows[0])
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        shade(c, NAVY)
        borders(c)
        cell_margin(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h))
        set_font(r, size=size, bold=True, color="FFFFFF")
        if widths:
            c.width = Cm(widths[j])
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        no_split(t.rows[-1])
        for j, value in enumerate(row):
            c = cells[j]
            shade(c, "FFFFFF" if i % 2 == 0 else LIGHT)
            borders(c)
            cell_margin(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            col = RED if isinstance(value, str) and value.startswith("(") else TEXT
            set_font(r, size=size, color=col)
            if widths:
                c.width = Cm(widths[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def add_figure(doc, path, caption, width=16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(6)
    r = cp.add_run(caption)
    set_font(r, size=8.8, italic=True, color="555555")


def chart_canvas(title, subtitle="Đơn vị: triệu đồng"):
    w, h = 1600, 800
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    d.text((105, 35), title, font=img_font(34, True), fill="#111827")
    if subtitle:
        d.text((105, 84), subtitle, font=img_font(19), fill="#6B7280")
    return im, d


def nice_bounds(values, include_zero=True):
    lo, hi = min(values), max(values)
    if include_zero:
        lo, hi = min(0, lo), max(0, hi)
    if hi == lo:
        hi = lo + 1
    span = hi - lo
    step0 = span / 5
    power = 10 ** max(0, len(str(int(abs(step0)))) - 1)
    step = ceil(step0 / power) * power
    lo2 = (lo // step) * step
    hi2 = ceil(hi / step) * step
    return float(lo2), float(hi2), float(step)


def line_chart(path, title, xlabels, series, subtitle="Đơn vị: triệu đồng", include_zero=True):
    im, d = chart_canvas(title, subtitle)
    left, top, right, bottom = 170, 145, 1510, 650
    allvals = [v for _, vals, _ in series for v in vals]
    lo, hi, step = nice_bounds(allvals, include_zero)
    y = lo
    while y <= hi + 1e-9:
        py = bottom - (y - lo) / (hi - lo) * (bottom - top)
        d.line((left, py, right, py), fill="#DCE3E8", width=2)
        label = fmt_m(y)
        box = d.textbbox((0, 0), label, font=img_font(17))
        d.text((left - 18 - box[2], py - 10), label, font=img_font(17), fill="#6B7280")
        y += step
    if lo < 0 < hi:
        py = bottom - (0 - lo) / (hi - lo) * (bottom - top)
        d.line((left, py, right, py), fill="#64748B", width=3)
    n = len(xlabels)
    xs = [left + (right - left) * i / max(1, n - 1) for i in range(n)]
    for i, x in enumerate(xs):
        label = str(xlabels[i])
        box = d.textbbox((0, 0), label, font=img_font(17))
        d.text((x - (box[2] - box[0]) / 2, bottom + 22), label, font=img_font(17), fill="#4B5563")
    lx = left
    for name, vals, color in series:
        pts = []
        for i, v in enumerate(vals):
            py = bottom - (v - lo) / (hi - lo) * (bottom - top)
            pts.append((xs[i], py))
        d.line(pts, fill=color, width=6, joint="curve")
        for x, y0 in pts:
            d.ellipse((x - 8, y0 - 8, x + 8, y0 + 8), fill=color, outline="white", width=2)
        d.line((lx, 720, lx + 42, 720), fill=color, width=6)
        d.text((lx + 54, 705), name, font=img_font(18), fill="#374151")
        lx += 330
    im.save(path, quality=95)


def grouped_bar(path, title, xlabels, series, subtitle="Đơn vị: triệu đồng"):
    im, d = chart_canvas(title, subtitle)
    left, top, right, bottom = 170, 145, 1510, 650
    vals = [v for _, arr, _ in series for v in arr]
    lo, hi, step = nice_bounds(vals, True)
    y = lo
    while y <= hi + 1e-9:
        py = bottom - (y - lo) / (hi - lo) * (bottom - top)
        d.line((left, py, right, py), fill="#DCE3E8", width=2)
        label = fmt_m(y)
        box = d.textbbox((0, 0), label, font=img_font(17))
        d.text((left - 18 - box[2], py - 10), label, font=img_font(17), fill="#6B7280")
        y += step
    py0 = bottom - (0 - lo) / (hi - lo) * (bottom - top)
    d.line((left, py0, right, py0), fill="#64748B", width=3)
    n, m = len(xlabels), len(series)
    group = (right - left) / n
    bw = group * 0.68 / m
    for i, lab in enumerate(xlabels):
        cx = left + group * (i + .5)
        box = d.textbbox((0, 0), str(lab), font=img_font(17))
        d.text((cx - (box[2]-box[0])/2, bottom + 22), str(lab), font=img_font(17), fill="#4B5563")
        for j, (_, arr, color) in enumerate(series):
            v = arr[i]
            x0 = cx - (m*bw)/2 + j*bw
            x1 = x0 + bw - 4
            py = bottom - (v - lo) / (hi - lo) * (bottom - top)
            d.rectangle((x0, min(py, py0), x1, max(py, py0)), fill=color)
    lx = left
    for name, _, color in series:
        d.rectangle((lx, 705, lx+35, 740), fill=color)
        d.text((lx+48, 707), name, font=img_font(18), fill="#374151")
        lx += 340
    im.save(path, quality=95)


def stacked_bar(path, title, xlabels, series, subtitle="Đơn vị: triệu đồng"):
    im, d = chart_canvas(title, subtitle)
    left, top, right, bottom = 170, 145, 1510, 650
    totals = [sum(s[1][i] for s in series) for i in range(len(xlabels))]
    lo, hi, step = nice_bounds(totals, True)
    y = 0
    while y <= hi + 1e-9:
        py = bottom - y / hi * (bottom - top)
        d.line((left, py, right, py), fill="#DCE3E8", width=2)
        label = fmt_m(y)
        box = d.textbbox((0, 0), label, font=img_font(17))
        d.text((left - 18 - box[2], py - 10), label, font=img_font(17), fill="#6B7280")
        y += step
    n = len(xlabels)
    group = (right-left)/n
    bw = group*.55
    for i, lab in enumerate(xlabels):
        x0 = left + group*(i+.5) - bw/2
        x1 = x0 + bw
        acc = 0
        for _, arr, color in series:
            y0 = bottom - acc/hi*(bottom-top)
            acc += arr[i]
            y1 = bottom - acc/hi*(bottom-top)
            d.rectangle((x0, y1, x1, y0), fill=color, outline="white")
        box = d.textbbox((0,0), str(lab), font=img_font(17))
        d.text(((x0+x1)/2-(box[2]-box[0])/2,bottom+22),str(lab),font=img_font(17),fill="#4B5563")
    lx = left
    for name, _, color in series:
        d.rectangle((lx, 705, lx+32, 737), fill=color)
        d.text((lx+45, 705), name, font=img_font(17), fill="#374151")
        lx += 285
    im.save(path, quality=95)


def donut(path, title, labels, values, colors):
    im, d = chart_canvas(title, "Cơ cấu theo số liệu trong file Excel")
    box = (120, 155, 700, 735)
    start = -90
    total = sum(values)
    for v, color in zip(values, colors):
        end = start + v/total*360
        d.pieslice(box, start=start, end=end, fill=color, outline="white", width=4)
        start = end
    d.ellipse((285, 320, 535, 570), fill="white")
    center = fmt_m(total)
    bb = d.textbbox((0,0),center,font=img_font(34,True))
    d.text((410-(bb[2]-bb[0])/2,397),center,font=img_font(34,True),fill="#163A5F")
    d.text((360,447),"triệu đồng",font=img_font(18),fill="#6B7280")
    y=185
    for lab,v,color in zip(labels,values,colors):
        d.rounded_rectangle((800,y,842,y+42),radius=6,fill=color)
        d.text((865,y-1),f"{lab}: {fmt_m(v)} ({v/total*100:.1f}%)",font=img_font(21),fill="#1F2937")
        y+=88
    im.save(path, quality=95)


# Workbook values (cached results produced by Excel).
wb = load_workbook(XLSX, data_only=True, read_only=False)
ass = wb["1. Giả định & Thông số"]
cash = wb["2. Báo cáo Ngân lưu"]
cost = wb["3. Chi phí chi tiết"]
debt = wb["5. Lịch trả nợ"]
income = wb["4. Báo cáo Thu nhập"]
evals = wb["6. Đánh Giá Tài Chính"]
infl = wb["7. Lạm phát & Khấu hao"]

years = [f"Năm {i}" for i in range(1, 8)]
partners = [cash.cell(6,c).value for c in range(3,10)]
saas = [cash.cell(5,c).value or 0 for c in range(3,10)]
transaction = [cash.cell(11,c).value or 0 for c in range(3,10)]
operating_revenue = [saas[i]+transaction[i] for i in range(7)]
nwc_recovery = [0]*6 + [cash.cell(14,9).value or 0]
total_inflow = [cash.cell(15,c).value for c in range(3,10)]
opex = [cost.cell(27,c).value for c in range(3,10)]
personnel = [cost.cell(10,c).value for c in range(3,10)]
technology = [cost.cell(15,c).value for c in range(3,10)]
marketing_rd = [cost.cell(20,c).value for c in range(3,10)]
fixed = [cost.cell(25,c).value for c in range(3,10)]
ebitda = [infl.cell(53,c).value for c in range(4,11)]
net_income = [income.cell(13,c).value for c in range(3,10)]
ebitda_margin = [infl.cell(55,c).value for c in range(4,11)]
net_margin = [income.cell(14,c).value for c in range(3,10)]
epv = [cash.cell(54,c).value for c in range(3,10)]
epv_cum = [cash.cell(55,c).value for c in range(3,10)]
discounted = [cash.cell(59,c).value for c in range(3,10)]
debt_interest = [debt.cell(5,c).value or 0 for c in range(3,10)]
debt_principal = [debt.cell(6,c).value or 0 for c in range(3,10)]
debt_end = [debt.cell(8,c).value or 0 for c in range(3,10)]
real_revenue = [infl.cell(20,c).value for c in range(4,11)]
real_ebitda = [infl.cell(54,c).value for c in range(4,11)]

initial_costs = [ass["B4"].value, ass["B5"].value, ass["B6"].value, ass["B7"].value]
capital_values = [ass["B13"].value, ass["B12"].value]

# Derived, directly reconciled metrics.
max_cum_deficit = abs(min([-ass["B13"].value] + epv_cum))
operating_bridge = max_cum_deficit - ass["B13"].value
total_funding_envelope = ass["B8"].value + operating_bridge
dpp = 5 + abs(sum([cash.cell(59,c).value for c in range(2,8)])) / cash["H59"].value
break_even_revenue = []
for i in range(7):
    fixed_opex = personnel[i] + 270 + 280
    break_even_revenue.append(fixed_opex / (1 - 0.37))

# Charts.
chart_capex = TMP / "01_capex.png"
donut(chart_capex, "Vốn đầu tư ban đầu: 2.590 triệu đồng",
      ["Thành lập", "Phát triển ứng dụng", "Thiết bị", "Triển khai"],
      initial_costs, ["#8AA6B8", "#1976A3", "#5FA777", "#D6A84B"])

chart_capital = TMP / "02_capital_mix.png"
donut(chart_capital, "Cấu trúc vốn trong mô hình",
      ["Vốn chủ sở hữu", "Vốn vay"], capital_values, ["#1976A3", "#D9822B"])

chart_revenue_mix = TMP / "03_revenue_mix.png"
stacked_bar(chart_revenue_mix, "Doanh thu vận hành theo nguồn", years,
            [("Phí sử dụng app", saas, "#1976A3"), ("Hoa hồng DN và ngân hàng", transaction, "#5FA777")])

chart_rev_opex = TMP / "04_revenue_opex.png"
line_chart(chart_rev_opex, "Doanh thu vận hành và chi phí vận hành", years,
           [("Doanh thu vận hành", operating_revenue, "#1976A3"), ("Chi phí vận hành", opex, "#D9822B")])

chart_profit = TMP / "05_profit.png"
grouped_bar(chart_profit, "EBITDA và lợi nhuận ròng", years,
            [("EBITDA", ebitda, "#24845B"), ("Lợi nhuận ròng theo Excel", net_income, "#7251A3")])

chart_partners = TMP / "06_partners.png"
line_chart(chart_partners, "Quy mô doanh nghiệp đối tác", years,
           [("Số doanh nghiệp", partners, "#1976A3")], subtitle="Đơn vị: doanh nghiệp", include_zero=True)

chart_margins = TMP / "07_margins.png"
line_chart(chart_margins, "Biên EBITDA và biên lợi nhuận ròng", years,
           [("Biên EBITDA", [x*100 for x in ebitda_margin], "#24845B"),
            ("Biên lợi nhuận ròng", [x*100 for x in net_margin], "#7251A3")],
           subtitle="Đơn vị: %", include_zero=True)

chart_breakeven = TMP / "08_breakeven.png"
line_chart(chart_breakeven, "Doanh thu thực tế so với ngưỡng hòa vốn EBITDA", years,
           [("Doanh thu vận hành", operating_revenue, "#1976A3"),
            ("Ngưỡng hòa vốn EBITDA", break_even_revenue, "#C73E3A")])

chart_cash = TMP / "09_cashflow.png"
line_chart(chart_cash, "Ngân lưu vốn chủ sở hữu: năm và lũy kế", years,
           [("NCF EPV từng năm", epv, "#24845B"), ("NCF EPV lũy kế", epv_cum, "#C73E3A")])

chart_debt = TMP / "10_debt.png"
line_chart(chart_debt, "Trả nợ và dư nợ cuối kỳ", years,
           [("Gốc", debt_principal, "#1976A3"), ("Lãi", debt_interest, "#D9822B"), ("Dư nợ cuối kỳ", debt_end, "#7251A3")])

chart_wacc = TMP / "11_wacc.png"
wacc_levels = [evals.cell(r,1).value*100 for r in range(24,29)]
npvs = [evals.cell(r,3).value for r in range(24,29)]
line_chart(chart_wacc, "Độ nhạy NPV theo chi phí vốn", [f"{x:.2f}%" for x in wacc_levels],
           [("NPV", npvs, "#1976A3")], include_zero=True)

chart_real = TMP / "12_real_nominal.png"
line_chart(chart_real, "Doanh thu danh nghĩa và doanh thu thực sau lạm phát", years,
           [("Danh nghĩa", operating_revenue, "#1976A3"), ("Thực sau CPI", real_revenue, "#D9822B")])

chart_cost = TMP / "13_cost_mix.png"
stacked_bar(chart_cost, "Cơ cấu chi phí vận hành", years,
            [("Nhân sự", personnel, "#1976A3"), ("Công nghệ", technology, "#24845B"),
             ("Marketing và R&D", marketing_rd, "#D9822B"), ("Cố định", fixed, "#7251A3")])


# Word document.
doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.7)
sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(1.8)
sec.right_margin = Cm(1.8)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(TEXT)
for name, size in (("Title", 22), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5)):
    st = doc.styles[name]
    st.font.name = FONT
    st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string("000000")
    st.paragraph_format.space_before = Pt(9)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PHẦN 5 - DỰ KIẾN TÀI CHÍNH")
set_font(r, size=22, bold=True, color="000000")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Dự án Champ Mạnh Chat | Kế hoạch tài chính 7 năm")
set_font(r, size=11.5, color="555555")

add_p(doc,
      "Phần tài chính này sử dụng trực tiếp các số liệu, công thức và kết quả đã lưu trong file BÁO CÁO TÀI CHÍNH SUZ2026 (4).xlsx. Không bổ sung giả định tăng trưởng mới. Cách trình bày tách riêng số đầu vào, kết quả tính toán và điều kiện thực thi để hội đồng có thể kiểm tra từng con số.",
      after=9)

doc.add_heading("5.1 Tóm tắt điều hành", level=2)
summary_rows = [
    ("Vốn đầu tư ban đầu", f"{fmt_m(ass['B8'].value)} triệu đồng", "Chi phí xây dựng tài sản và triển khai ban đầu"),
    ("Cơ cấu vốn trong mô hình", f"{fmt_m(ass['B13'].value)} triệu vốn chủ; {fmt_m(ass['B12'].value)} triệu vốn vay", "60% vốn chủ - 40% vốn vay"),
    ("Doanh thu vận hành", f"{fmt_m(operating_revenue[0])} triệu Năm 1; {fmt_m(operating_revenue[-1])} triệu Năm 7", "Không tính thu hồi vốn lưu động là doanh thu"),
    ("Hòa vốn EBITDA", "Năm 3", f"EBITDA {fmt_m(ebitda[2])} triệu; biên {fmt_pct(ebitda_margin[2])}"),
    ("Lợi nhuận ròng dương", "Năm 4", f"{fmt_m(net_income[3])} triệu đồng theo báo cáo thu nhập"),
    ("Hoàn vốn không chiết khấu", "5,16 năm", "Dòng tiền vốn chủ lũy kế dương trong Năm 6"),
    ("Hoàn vốn chiết khấu", f"Khoảng {dpp:.2f} năm".replace(".", ","), "Tính lại từ dòng NCF EPV đã chiết khấu trong Excel"),
    ("NPV / IRR / MIRR", f"{fmt_m(evals['B5'].value)} triệu / {fmt_pct(evals['B6'].value,2)} / {fmt_pct(evals['B7'].value,2)}", f"So với WACC {fmt_pct(ass['B17'].value,2)}"),
    ("Khoảng vốn vận hành phải khóa thêm", f"{fmt_m(operating_bridge)} triệu đồng", "Để bù dòng tiền âm đến cuối Năm 3"),
]
add_table(doc, ["Chỉ tiêu", "Kết quả", "Cách hiểu thực tế"], summary_rows, widths=[4.8,5.3,7.0], size=8.8)

add_p(doc,
      "Kết luận ngắn: mô hình có khả năng tạo lợi nhuận và giá trị nếu đạt đủ quy mô doanh nghiệp đối tác, nhưng 2.590 triệu đồng vốn đầu tư ban đầu chưa phải toàn bộ số tiền cần huy động. Dòng tiền hoạt động âm trong ba năm đầu làm phát sinh nhu cầu vốn bổ sung khoảng 2.402 triệu đồng. Vì vậy, điều kiện sống còn không chỉ là đạt doanh thu mà còn là khóa đủ runway trước khi triển khai toàn bộ chi phí ứng dụng và vận hành.",
      bold_lead="Kết luận ngắn:")

doc.add_heading("5.2 Vốn đầu tư ban đầu và cấu trúc tài trợ", level=2)
capex_rows = [
    ("Thành lập doanh nghiệp", fmt_m(initial_costs[0]), fmt_pct(initial_costs[0]/ass['B8'].value)),
    ("Phát triển ứng dụng", fmt_m(initial_costs[1]), fmt_pct(initial_costs[1]/ass['B8'].value)),
    ("Vật tư và thiết bị", fmt_m(initial_costs[2]), fmt_pct(initial_costs[2]/ass['B8'].value)),
    ("Ra mắt và triển khai", fmt_m(initial_costs[3]), fmt_pct(initial_costs[3]/ass['B8'].value)),
    ("Tổng", fmt_m(ass['B8'].value), "100,0%"),
]
add_table(doc, ["Hạng mục", "Triệu đồng", "Tỷ trọng"], capex_rows, widths=[8.2,4.0,3.5], size=9.2)
add_figure(doc, chart_capex, "Hình 1. Phân bổ vốn đầu tư ban đầu theo file tài chính")

add_p(doc,
      "Chi phí phát triển ứng dụng chiếm 69,5% tổng đầu tư. Đây là khoản chi lớn nhất và cũng là khoản có thể kiểm soát theo milestone. Cách triển khai thực tế là chia nghiệm thu thành MVP, dữ liệu và cảnh báo, giao dịch - uy tín đối tác, sau đó mới tích hợp ngân hàng/bảo hiểm. Không nên thanh toán toàn bộ 1.800 triệu đồng trước khi có sản phẩm chạy thử và biên bản nghiệm thu từng giai đoạn.")

capital_rows = [
    ("Vốn chủ sở hữu", fmt_m(ass['B13'].value), "60,0%", "Chịu phần rủi ro ban đầu"),
    ("Vốn vay", fmt_m(ass['B12'].value), "40,0%", f"Lãi suất {fmt_pct(ass['B14'].value)}; thời hạn {int(ass['B18'].value)} năm"),
    ("Tổng", fmt_m(ass['B8'].value), "100,0%", f"WACC {fmt_pct(ass['B17'].value,2)}"),
]
add_table(doc, ["Nguồn", "Triệu đồng", "Tỷ trọng", "Điều kiện"], capital_rows, widths=[4.3,3.4,3.0,6.0], size=9)
add_figure(doc, chart_capital, "Hình 2. Cơ cấu vốn chủ và vốn vay trong mô hình")

doc.add_heading("Khoảng vốn phải chuẩn bị ngoài 2.590 triệu đồng", level=3)
funding_rows = [
    ("Vốn đầu tư tài sản ban đầu", fmt_m(ass['B8'].value), "Ứng dụng, thiết bị, triển khai, thành lập"),
    ("Dòng tiền hoạt động phải bù đến cuối Năm 3", fmt_m(operating_bridge), "Chênh lệch âm sau doanh thu và trả nợ"),
    ("Tổng nguồn vốn tối thiểu cần có khả năng tiếp cận", fmt_m(total_funding_envelope), "Tương đương vốn đầu tư cộng buffer vận hành"),
]
add_table(doc, ["Cấu phần", "Triệu đồng", "Ý nghĩa"], funding_rows, widths=[7.0,3.7,6.2], size=9)
add_p(doc,
      "Đây là điểm quyết định tính khả thi. Nếu chỉ huy động 2.590 triệu đồng rồi sử dụng hết cho tài sản ban đầu, dự án sẽ thiếu tiền để trả nhân sự, duy trì hạ tầng và trả nợ trong ba năm đầu. Khoản vốn vận hành 2.402 triệu đồng phải được cam kết theo tranche, SAFE, grant, vốn đối tác chiến lược hoặc khoản vay có thời gian ân hạn phù hợp trước khi ký toàn bộ hợp đồng phát triển.")

doc.add_heading("5.3 Doanh thu hình thành từ hoạt động nào", level=2)
revenue_model_rows = [
    ("Phí sử dụng app", "5 triệu đồng/DN/tháng", "8 DN Năm 1 lên 220 DN Năm 7", "Doanh thu định kỳ"),
    ("Hoa hồng DN và ngân hàng", "5% giao dịch trả trước; công thức Excel dùng 6% cho trả sau", "Chỉ tính trên giao dịch đi qua nền tảng", "Doanh thu theo giao dịch"),
    ("Hoa hồng bảo hiểm", "Chưa ghi nhận", "Các ô doanh thu đang để trống", "Không đưa vào kết quả cơ sở"),
]
add_table(doc, ["Nguồn doanh thu", "Mức giá/tỷ lệ", "Cơ sở khối lượng", "Tính chất"], revenue_model_rows, widths=[4.2,4.1,5.1,3.6], size=8.8)

revenue_rows=[]
for i,y in enumerate(years):
    revenue_rows.append((y, int(partners[i]), fmt_m(saas[i]), fmt_m(transaction[i]), fmt_m(operating_revenue[i]), fmt_m(nwc_recovery[i]), fmt_m(total_inflow[i])))
add_table(doc, ["Năm", "DN", "Phí app", "Hoa hồng", "Doanh thu vận hành", "Thu hồi NWC", "Tổng tiền vào"], revenue_rows,
          widths=[1.8,1.5,2.4,2.5,3.3,2.4,2.6], size=7.9)
add_figure(doc, chart_revenue_mix, "Hình 3. Doanh thu vận hành theo hai nguồn được ghi nhận")

add_p(doc,
      "Năm 7 cần đọc đúng: doanh thu vận hành là 25.520 triệu đồng, gồm 13.200 triệu phí app và 12.320 triệu hoa hồng. Khoản 2.552 triệu đồng còn lại là thu hồi vốn lưu động cuối dự án, nên thuộc ngân lưu vào chứ không phải doanh thu bán hàng. Việc tách khoản này giúp chỉ tiêu doanh thu, biên lợi nhuận và khả năng scale không bị trình bày cao hơn thực tế.")

doc.add_heading("5.4 Doanh thu, chi phí và lợi nhuận 7 năm", level=2)
financial_rows=[]
for i,y in enumerate(years):
    financial_rows.append((y, fmt_m(operating_revenue[i]), fmt_m(opex[i]), fmt_m(ebitda[i]), fmt_pct(ebitda_margin[i]), fmt_m(net_income[i]), fmt_pct(net_margin[i])))
add_table(doc, ["Năm", "Doanh thu vận hành", "Opex", "EBITDA", "Biên EBITDA", "LN ròng theo Excel", "Biên LN ròng"],
          financial_rows, widths=[1.7,3.3,2.6,2.5,2.5,3.0,2.6], size=7.9)
add_figure(doc, chart_rev_opex, "Hình 4. Doanh thu vận hành và chi phí vận hành")
add_figure(doc, chart_profit, "Hình 5. EBITDA và lợi nhuận ròng theo báo cáo tài chính")
add_figure(doc, chart_margins, "Hình 6. Quá trình cải thiện biên EBITDA và biên lợi nhuận ròng")

add_p(doc,
      "Ba năm đầu là giai đoạn hấp thụ chi phí cố định. EBITDA lần lượt âm 897 triệu, âm 485 triệu và dương 114 triệu đồng. Đến Năm 4, lợi nhuận ròng chuyển dương 704 triệu đồng. Từ Năm 5 trở đi, tăng trưởng doanh thu nhanh hơn tăng chi phí cố định làm biên EBITDA mở rộng mạnh. Đây là cơ chế sinh lợi chính của mô hình SaaS kết hợp phí giao dịch.")

doc.add_heading("5.5 Chi phí được kiểm soát như thế nào", level=2)
cost_rows=[]
for i,y in enumerate(years):
    cost_rows.append((y, fmt_m(personnel[i]), fmt_m(technology[i]), fmt_m(marketing_rd[i]), fmt_m(fixed[i]), fmt_m(opex[i])))
add_table(doc, ["Năm", "Nhân sự", "Công nghệ", "Marketing và R&D", "Chi phí cố định", "Tổng Opex"], cost_rows,
          widths=[1.8,2.8,2.8,3.4,3.2,3.0], size=8.2)
add_figure(doc, chart_cost, "Hình 7. Cơ cấu chi phí vận hành theo nhóm")

add_p(doc,
      "Chi phí biến đổi trong công thức thực tế gồm 15% doanh thu cho hoạt động nền tảng, 12% cho marketing và 10% cho R&D, tổng cộng 37% doanh thu. Phần còn lại là nhân sự, duy trì công nghệ 270 triệu đồng/năm và chi phí quản lý - pháp lý 280 triệu đồng/năm. Vì vậy, mỗi 1 đồng doanh thu tăng thêm tạo khoảng 0,63 đồng đóng góp trước chi phí cố định, khấu hao, lãi vay và thuế.")

doc.add_heading("5.6 Điểm hòa vốn: ba mốc phải phân biệt", level=2)
breakeven_rows = [
    ("Hòa vốn EBITDA", "Năm 3", f"Doanh thu {fmt_m(operating_revenue[2])} triệu vượt ngưỡng khoảng {fmt_m(break_even_revenue[2])} triệu", f"EBITDA dương {fmt_m(ebitda[2])} triệu"),
    ("Lợi nhuận ròng dương", "Năm 4", "Sau khấu hao, lãi vay và thuế", f"Lợi nhuận ròng {fmt_m(net_income[3])} triệu"),
    ("Hoàn vốn không chiết khấu", "5,16 năm", "NCF EPV lũy kế chuyển dương trong Năm 6", "Theo công thức PP trong Excel"),
    ("Hoàn vốn chiết khấu", f"Khoảng {dpp:.2f} năm".replace(".", ","), "Tính từ NCF EPV đã chiết khấu", "Thận trọng hơn PP"),
]
add_table(doc, ["Mốc", "Thời điểm", "Cách tính", "Kết quả"], breakeven_rows, widths=[4.1,2.5,6.3,4.0], size=8.8)
add_figure(doc, chart_breakeven, "Hình 8. Doanh thu vận hành và ngưỡng hòa vốn EBITDA")

add_p(doc,
      "Mốc dùng để điều hành là hòa vốn EBITDA, vì nó trả lời hoạt động cốt lõi đã tự nuôi được bộ máy hay chưa. Mốc lợi nhuận ròng dùng cho báo cáo kế toán. Mốc hoàn vốn trả lời nhà đầu tư phải chờ bao lâu để dòng tiền lũy kế bù lại số vốn đã bỏ ra. Không nên dùng ba khái niệm này thay thế cho nhau.")

doc.add_heading("5.7 Dòng tiền và thời gian hoàn vốn", level=2)
cash_rows=[]
for i,y in enumerate(years):
    cash_rows.append((y, fmt_m(epv[i]), fmt_m(epv_cum[i]), fmt_m(discounted[i])))
add_table(doc, ["Năm", "NCF vốn chủ", "NCF vốn chủ lũy kế", "NCF đã chiết khấu"], cash_rows,
          widths=[2.5,4.4,4.8,4.4], size=9)
add_figure(doc, chart_cash, "Hình 9. Ngân lưu vốn chủ sở hữu từng năm và lũy kế")

add_p(doc,
      "Dòng tiền vốn chủ âm 1.237 triệu đồng ở Năm 1, âm 847 triệu ở Năm 2 và âm 318 triệu ở Năm 3. Mức âm lũy kế lớn nhất là 3.956 triệu đồng vào cuối Năm 3, bao gồm 1.554 triệu đồng vốn chủ ban đầu. Từ Năm 4, dự án bắt đầu tạo tiền dương; đến Năm 6, dòng tiền lũy kế vượt điểm hoàn vốn.")

doc.add_heading("5.8 Khả năng trả nợ", level=2)
debt_rows=[]
for i,y in enumerate(years):
    debt_rows.append((y, fmt_m(debt_interest[i]), fmt_m(debt_principal[i]), fmt_m(debt_interest[i]+debt_principal[i]), fmt_m(debt_end[i])))
add_table(doc, ["Năm", "Lãi", "Gốc", "Tổng trả nợ", "Dư nợ cuối kỳ"], debt_rows,
          widths=[2.3,3.0,3.0,3.8,3.8], size=9)
add_figure(doc, chart_debt, "Hình 10. Nghĩa vụ trả nợ và dư nợ cuối kỳ")

add_p(doc,
      "Khoản vay 1.036 triệu đồng tạo nghĩa vụ trả 280,31 triệu đồng mỗi năm trong 5 năm. Trong khi đó, dòng tiền trước tài trợ vẫn âm ở Năm 1, Năm 2 và Năm 3. Vì vậy, khoản vay thương mại chỉ khả thi nếu có tài sản bảo đảm, nguồn trả nợ khác, ân hạn gốc hoặc đối tác cam kết dòng tiền. Phương án thực thi an toàn hơn là giải ngân nợ theo milestone, ưu tiên grant/SAFE/vốn đối tác cho phần rủi ro cao và chỉ tăng nợ khi doanh thu định kỳ đủ chứng minh khả năng trả nợ.")

doc.add_heading("5.9 NPV, IRR và hiệu quả vốn", level=2)
returns_rows = [
    ("WACC", fmt_pct(ass['B17'].value,2), "Mức chiết khấu của mô hình"),
    ("NPV", f"{fmt_m(evals['B5'].value)} triệu đồng", "Giá trị hiện tại ròng dương"),
    ("IRR", fmt_pct(evals['B6'].value,2), f"Cao hơn WACC {fmt_pct(ass['B17'].value,2)}"),
    ("MIRR", fmt_pct(evals['B7'].value,2), "Thận trọng hơn IRR vì tái đầu tư theo WACC"),
    ("BCR", f"{evals['B8'].value:.2f} lần".replace(".", ","), "Giá trị chiết khấu tạo ra trên vốn chủ ban đầu"),
    ("EAB", f"{fmt_m(evals['B9'].value)} triệu đồng/năm", "NPV quy đổi thành lợi ích đều hàng năm"),
]
add_table(doc, ["Chỉ tiêu", "Kết quả", "Cách đọc"], returns_rows, widths=[4.0,4.5,8.0], size=9.1)
add_figure(doc, chart_wacc, "Hình 11. NPV vẫn dương trong dải WACC 10%-25% của file Excel")

add_p(doc,
      "NPV 5.110 triệu đồng và IRR 34,95% cho thấy mô hình có biên sinh lợi trên chi phí vốn 14,32%. Tuy nhiên, đây là kết quả của mô hình 7 năm chứ không phải lợi nhuận đã ghi nhận. Hai biến quyết định là tốc độ tăng doanh nghiệp trả phí và tỷ lệ giao dịch thực sự đi qua nền tảng. Sau mỗi quý pilot, nhóm phải thay số lượng DN, GMV và chi phí thực tế vào file trước khi tiếp tục dùng NPV/IRR để gọi vốn.")

doc.add_heading("5.10 Khả năng scale", level=2)
scale_rows=[]
for i,y in enumerate(years):
    revenue_per_partner = operating_revenue[i]/partners[i]
    scale_rows.append((y, int(partners[i]), fmt_m(operating_revenue[i]), fmt_m(revenue_per_partner,1), fmt_pct(ebitda_margin[i])))
add_table(doc, ["Năm", "DN đối tác", "Doanh thu vận hành", "Doanh thu/DN", "Biên EBITDA"], scale_rows,
          widths=[2.1,2.8,4.2,3.7,3.2], size=9)
add_figure(doc, chart_partners, "Hình 12. Tăng trưởng số doanh nghiệp đối tác")

add_p(doc,
      "Số doanh nghiệp tăng từ 8 lên 220, tương đương CAGR 73,7% theo file. Doanh thu vận hành trên mỗi doanh nghiệp tăng từ khoảng 74 triệu lên 116 triệu đồng/năm nhờ hoa hồng giao dịch tăng bên cạnh phí thuê bao cố định. Đây là cơ sở để biên EBITDA tăng từ âm lên 52,7% ở Năm 7. Khả năng scale chỉ được xác nhận khi chi phí cloud, hỗ trợ và xử lý giao dịch trên mỗi DN không tăng nhanh hơn doanh thu trên mỗi DN.")

doc.add_heading("5.11 Ảnh hưởng của lạm phát", level=2)
inflation_rows=[]
for i,y in enumerate(years):
    inflation_rows.append((y, fmt_m(operating_revenue[i]), fmt_m(real_revenue[i]), fmt_m(ebitda[i]), fmt_m(real_ebitda[i])))
add_table(doc, ["Năm", "Doanh thu danh nghĩa", "Doanh thu thực", "EBITDA danh nghĩa", "EBITDA thực"], inflation_rows,
          widths=[2.1,3.6,3.3,3.6,3.3], size=8.8)
add_figure(doc, chart_real, "Hình 13. Chênh lệch doanh thu danh nghĩa và doanh thu thực sau CPI 4,5%")

add_p(doc,
      "File sử dụng CPI 4,5%/năm. Doanh thu Năm 7 là 25.520 triệu đồng theo giá danh nghĩa nhưng sức mua thực thấp hơn. Vì vậy, quyết định tuyển dụng, thuê hạ tầng và giá gói dịch vụ phải dựa trên cả biên danh nghĩa và biên thực, tránh nhầm tăng giá tiền tệ với tăng hiệu quả kinh doanh.")

doc.add_heading("5.12 Các điểm phải khóa trước khi nộp hồ sơ hoặc gọi vốn", level=2)
audit_rows = [
    ("Doanh thu Năm 7", "Tách 25.520 triệu doanh thu vận hành khỏi 2.552 triệu thu hồi vốn lưu động", "Sửa đồng thời công thức Opex, thuế và lợi nhuận Năm 7 để không dùng khoản thu hồi làm doanh thu"),
    ("Tỷ lệ giao dịch trả sau", "Ô giả định ghi 5,95% nhưng công thức doanh thu đang dùng 6%", "Chọn một tỷ lệ và cập nhật toàn bộ mô hình"),
    ("Tăng lương", "Ô giả định là 5% nhưng ghi chú tại một số sheet còn 7%", "Khóa một tỷ lệ chính thức"),
    ("Tỷ lệ Marketing và R&D", "Công thức dùng 12% và 10% nhưng nhãn tại bảng ngân lưu còn ghi 15% và 12%", "Đồng bộ nhãn với công thức 12% và 10%"),
    ("Vốn vay", "Công thức và lịch trả nợ dùng 1.036 triệu nhưng dòng mô tả trên sheet còn ghi 957 triệu", "Sửa dòng mô tả thành 1.036 triệu"),
    ("Chi phí vốn chủ", "Ô giả định dùng 18% nhưng ghi chú WACC tại sheet ngân lưu còn ghi Re=20%", "Giữ Re=18% nếu tiếp tục dùng WACC 14,32%"),
    ("Giai đoạn 7 năm", "Một sheet ghi 2026-2034; bảy năm liên tục phải là 2026-2032", "Dùng Năm 1-Năm 7 hoặc sửa niên độ"),
    ("Phân tích doanh thu nhạy cảm", "Bảng NPV theo 50%-150% doanh thu trong Excel không tăng đơn điệu", "Không sử dụng bảng này cho tới khi sửa công thức"),
    ("Nguồn vốn vận hành", f"Phải chứng minh khả năng tiếp cận thêm {fmt_m(operating_bridge)} triệu", "Không triển khai toàn bộ capex nếu chưa khóa runway"),
]
add_table(doc, ["Vấn đề", "Hiện trạng", "Hành động"], audit_rows, widths=[4.0,7.0,6.0], size=8.7)

doc.add_heading("5.13 Cơ chế kiểm soát để con số biến thành kết quả thực", level=2)
controls = [
    "Giải ngân phát triển ứng dụng theo nghiệm thu: MVP, dữ liệu/cảnh báo, giao dịch/uy tín, tích hợp đối tác.",
    "Mỗi tháng theo dõi bốn số bắt buộc: DN đang trả phí, doanh thu định kỳ, GMV hoàn tất và tiền cuối kỳ.",
    "Không tính lead, đăng ký hoặc bài đăng là giao dịch. Chỉ ghi nhận hoa hồng khi giao dịch hoàn tất và đối soát được.",
    "Dừng tuyển dụng và capex chưa thiết yếu nếu runway xuống dưới 9 tháng hoặc doanh thu đạt dưới 75% kế hoạch hai quý liên tiếp.",
    "Không tăng vay thương mại khi EBITDA và dòng tiền trả nợ chưa dương; ưu tiên nguồn vốn chia sẻ rủi ro ở giai đoạn đầu.",
    "Cập nhật lại NPV/IRR hàng quý bằng dữ liệu đã phát sinh, không dùng cố định bộ giả định ban đầu để thuyết phục nhà đầu tư.",
]
for item in controls:
    add_bullet(doc, item)

doc.add_heading("5.14 Kết luận tài chính", level=2)
add_p(doc,
      "Mô hình cho thấy Champ Mạnh Chat có thể đạt hòa vốn EBITDA ở Năm 3, lợi nhuận ròng dương ở Năm 4 và hoàn vốn khoảng 5,16 năm nếu tăng trưởng đúng số doanh nghiệp và giao dịch trong file. Giá trị tài chính nằm ở việc kết hợp doanh thu thuê bao định kỳ với phí giao dịch, trong khi chi phí cố định được hấp thụ dần khi quy mô tăng. NPV dương 5.110 triệu đồng, IRR 34,95% và MIRR 30,02% cho thấy dư địa sinh lợi trên WACC 14,32%.")
add_p(doc,
      "Điểm phải xử lý trước khi triển khai là vốn. Tổng đầu tư tài sản 2.590 triệu đồng chưa bao gồm đủ buffer vận hành; dự án cần khóa thêm khoảng 2.402 triệu đồng để đi qua ba năm dòng tiền âm. Kế hoạch được đánh giá cao không phải vì con số lợi nhuận lớn nhất, mà vì đội thi nhìn thấy đúng khoảng thiếu tiền, tách đúng doanh thu khỏi thu hồi vốn lưu động và chỉ mở rộng khi khách hàng trả phí cùng dòng tiền thực xác nhận mô hình.")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
r = p.add_run("Nguồn số liệu: BÁO CÁO TÀI CHÍNH SUZ2026 (4).xlsx. Đơn vị trong tài liệu là triệu đồng, trừ khi ghi chú khác. Các chỉ tiêu NPV, IRR, MIRR, lợi nhuận và hoàn vốn là kết quả tính toán của mô hình tài chính 7 năm, không phải doanh thu hoặc lợi nhuận đã thực hiện.")
set_font(r, size=9, italic=True, color="555555")

for p in doc.paragraphs:
    ppr = p._p.get_or_add_pPr()
    wid = OxmlElement("w:widowControl")
    wid.set(qn("w:val"), "1")
    ppr.append(wid)

doc.core_properties.title = "Phần 5 - Dự kiến tài chính final"
doc.core_properties.subject = "Dự án Champ Mạnh Chat - Startup Zone 2026"
doc.core_properties.author = "Nhóm dự án Champ Mạnh Chat"
doc.save(OUT)
print(OUT)
