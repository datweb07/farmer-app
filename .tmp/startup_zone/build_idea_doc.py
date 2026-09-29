from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path(r"D:\final app\final app\docs\[STARTUP ZONE 2026] YÊU CẦU NỘI DUNG FILE Ý TƯỞNG - CHAMP MẠNH CHAT.docx")

GREEN = "1F5D42"
GREEN_2 = "2F7D57"
LIGHT_GREEN = "EAF4EE"
PALE = "F6F8F7"
GOLD = "D6A21E"
DARK = "1F2937"
MID = "52606D"
WHITE = "FFFFFF"
GRAY = "D8DEE4"
RED = "A63D40"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRAY, size="5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color=DARK, size=9, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False if widths else True
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, bold=True, color=WHITE, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(hdr.cells[i], GREEN)
        set_cell_border(hdr.cells[i], GREEN)
        hdr.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            hdr.cells[i].width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cells[i], value, size=font_size, align=align)
            set_cell_shading(cells[i], WHITE if ridx % 2 == 0 else PALE)
            set_cell_border(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[i].width = Inches(widths[i])
    return table


def add_p(doc, text="", style=None, bold_prefix=None, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_together = keep
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def number(doc, text, level=0):
    p = doc.add_paragraph(style="List Number" if level == 0 else "List Number 2")
    p.add_run(text)
    return p


def page_break(doc):
    doc.add_page_break()


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.62)
sec.bottom_margin = Inches(0.62)
sec.left_margin = Inches(0.7)
sec.right_margin = Inches(0.7)
sec.header_distance = Inches(0.28)
sec.footer_distance = Inches(0.28)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.2)
styles["Normal"].font.color.rgb = RGBColor.from_string(DARK)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.08

for name, size, color, before, after in [
    ("Title", 28, GREEN, 0, 8),
    ("Heading 1", 17, GREEN, 16, 7),
    ("Heading 2", 12.5, GREEN_2, 11, 5),
    ("Heading 3", 10.5, DARK, 8, 3),
]:
    s = styles[name]
    s.font.name = "Aptos Display" if name != "Normal" else "Aptos"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

if "Eyebrow" not in styles:
    s = styles.add_style("Eyebrow", WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = "Aptos"
    s.font.size = Pt(9)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(GOLD)
    s.paragraph_format.space_after = Pt(8)

header = sec.header
hp = header.paragraphs[0]
hp.text = "STARTUP ZONE 2026 · BẢNG A"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(GREEN)

footer = sec.footer
fp = footer.paragraphs[0]
fp.text = "CHAMP MẠNH CHAT · FILE Ý TƯỞNG"
for r in fp.runs:
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MID)
add_page_number(footer.add_paragraph())

# Cover
p = doc.add_paragraph(style="Eyebrow")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("STARTUP ZONE LẦN 10 · TECTONIC · NĂM 2026")

spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(48)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("FILE Ý TƯỞNG")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(16)
r.font.color.rgb = RGBColor.from_string(GOLD)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("CHAMP MẠNH CHAT")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Nền tảng kết nối Nông dân – Doanh nghiệp – Ngân hàng chính sách\ncho nông nghiệp thích ứng khí hậu Đồng bằng sông Cửu Long")
r.font.size = Pt(14)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string(DARK)

t = doc.add_table(rows=1, cols=1)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
c = t.cell(0, 0)
set_cell_shading(c, LIGHT_GREEN)
set_cell_border(c, GREEN, "10")
set_cell_text(c, "DỮ LIỆU → QUYẾT ĐỊNH → NGUỒN LỰC → KẾT QUẢ → DỮ LIỆU THỰC TẾ", bold=True, color=GREEN, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_paragraph()
add_table(doc, ["Hạng mục", "Thông tin"], [
    ("Bảng dự thi", "Bảng A – Sinh viên khởi nghiệp"),
    ("Lĩnh vực chính", "Công nghệ & chuyển đổi số; Nông nghiệp & thực phẩm"),
    ("Giai đoạn", "MVP đã xây dựng · đang hoàn thiện và chuẩn bị pilot"),
    ("Địa bàn khởi đầu", "Đồng bằng sông Cửu Long; ưu tiên mô hình sầu riêng tại Cần Thơ"),
    ("Đại diện đội", "Phan Khắc Anh Tuấn · UEH · 0585 708 372"),
    ("Email đội", "khacanhtuanphan@gmail.com"),
], widths=[1.7, 5.4], font_size=9.3)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("“Không dừng ở cảnh báo rủi ro; CHAMP MẠNH CHAT giúp người nông dân biết bước tiếp theo cần làm, kết nối đúng nguồn lực và tạo ra dữ liệu đáng tin cho mùa vụ kế tiếp.”")
r.italic = True
r.font.size = Pt(10.5)
r.font.color.rgb = RGBColor.from_string(GREEN)

page_break(doc)

# Executive summary
doc.add_heading("TÓM TẮT ĐỀ XUẤT", level=1)
add_p(doc, "CHAMP MẠNH CHAT là nền tảng số ba bên dành cho nông nghiệp thích ứng khí hậu ở Đồng bằng sông Cửu Long (ĐBSCL). Dự án kết nối nông dân, doanh nghiệp thu mua/cung ứng và tổ chức tài chính chính sách trên cùng một dòng dữ liệu: rủi ro độ mặn – nhu cầu sản xuất – giao dịch – mức độ uy tín – kết quả thực tế.")
add_p(doc, "Điểm khác biệt không nằm ở việc có thêm một ứng dụng cảnh báo hay một chợ nông sản. CHAMP MẠNH CHAT biến dữ liệu khí hậu thành hành động có thể thực hiện: nông dân nhận cảnh báo và khuyến nghị; doanh nghiệp nhìn thấy nhu cầu, vùng nguyên liệu và lịch thu mua; tổ chức tài chính có thêm dữ liệu phục vụ thiết kế và theo dõi nguồn vốn. Mỗi giao dịch hoàn tất lại tạo dữ liệu để tăng độ tin cậy cho quyết định tiếp theo.")

add_table(doc, ["Giá trị cho nông dân", "Giá trị cho doanh nghiệp", "Giá trị cho đối tác tài chính"], [
    ("Cảnh báo độ mặn theo địa bàn; thông tin dễ hiểu; tiếp cận đầu vào, đầu ra và nguồn lực phù hợp.",
     "Tìm nguồn cung theo khu vực; đăng nhu cầu thu mua; theo dõi đăng ký; xây dựng điểm uy tín từ giao dịch và đánh giá xác thực.",
     "Có thêm dữ liệu mùa vụ, giao dịch và rủi ro để sàng lọc, giám sát và thiết kế sản phẩm phù hợp hơn."),
], widths=[2.35, 2.35, 2.35], font_size=8.8)

doc.add_heading("Luận điểm đầu tư/đồng hành", level=2)
bullet(doc, "Vấn đề đủ lớn và cấp thiết: xâm nhập mặn tác động trực tiếp đến thời điểm tưới, lựa chọn đầu vào, năng suất và khả năng trả nợ của nông hộ.")
bullet(doc, "MVP đã hiện thực hóa nhiều mắt xích cốt lõi: bản đồ và dữ liệu dự báo độ mặn, cảnh báo, hồ sơ/GPS, cộng đồng, đăng nhu cầu thu mua, đăng ký giao dịch và đánh giá uy tín doanh nghiệp.")
bullet(doc, "Mô hình doanh thu B2B giúp giữ khả năng tiếp cận cho nông dân: phí nền tảng từ doanh nghiệp và hoa hồng theo giao dịch/dịch vụ đối tác.")
bullet(doc, "Chiến lược triển khai có trọng tâm: bắt đầu bằng một cây trồng giá trị cao và một địa bàn đủ hẹp để đo kết quả, sau đó chuẩn hóa trước khi mở rộng.")

doc.add_heading("Các mốc tài chính theo mô hình kế hoạch", level=2)
add_table(doc, ["Đầu tư ban đầu", "Doanh thu Năm 1", "Doanh thu Năm 7", "IRR kế hoạch", "NPV kế hoạch"], [
    ("3,01 tỷ đồng", "1,58 tỷ đồng", "20,691 tỷ đồng", "36,11%", "4,442 tỷ đồng"),
], widths=[1.38]*5, font_size=9)
add_p(doc, "Lưu ý: đây là mô hình kế hoạch 7 năm theo file tài chính của đội, không phải cam kết kết quả. Các giả định về số doanh nghiệp trả phí, giá thuê bao và giá trị giao dịch sẽ được khóa lại sau pilot.", style=None)

page_break(doc)

# Part 1
doc.add_heading("PHẦN 1 · GIỚI THIỆU DỰ ÁN", level=1)
doc.add_heading("1.1. Thông tin chung", level=2)
add_table(doc, ["Nội dung", "Mô tả"], [
    ("Tên dự án", "CHAMP MẠNH CHAT"),
    ("Định vị", "Nền tảng hỗ trợ quyết định và kết nối nguồn lực cho nông nghiệp thích ứng khí hậu ĐBSCL."),
    ("Sản phẩm", "Ứng dụng web/mobile-responsive kết hợp bản đồ, dữ liệu dự báo, cảnh báo, cộng đồng, thị trường thu mua và hồ sơ uy tín."),
    ("Khách hàng trả phí chính", "Doanh nghiệp thu mua, doanh nghiệp đầu vào và đối tác tổ chức; nông dân là người dùng hưởng lợi trực tiếp."),
    ("Giai đoạn", "MVP phát triển nội bộ; chưa thương mại hóa; đang hoàn thiện, kiểm chứng mô hình và chuẩn bị pilot."),
], widths=[1.65, 5.45], font_size=9)

doc.add_heading("1.2. Sản phẩm và tính năng", level=2)
add_table(doc, ["Mô-đun", "Giá trị cung cấp", "Trạng thái hiện tại"], [
    ("Theo dõi & dự báo độ mặn", "Bản đồ điểm quan trắc, biểu đồ, lọc dữ liệu, dự báo Prophet và cảnh báo theo ngưỡng.", "Đã có trong MVP"),
    ("Hồ sơ nông hộ & vị trí", "Khu vực sinh sống, vị trí GPS, hồ sơ cá nhân và dữ liệu phục vụ cá nhân hóa.", "Đã có trong MVP"),
    ("Cộng đồng tri thức", "Bài viết, bình luận, thích, chia sẻ, theo dõi, thông báo và kiểm duyệt nội dung.", "Đã có trong MVP"),
    ("Thu mua nông sản", "Doanh nghiệp đăng nhu cầu; nông dân đăng ký; doanh nghiệp xác nhận, cập nhật khối lượng và hoàn tất giao dịch.", "Đã có luồng MVP"),
    ("Uy tín doanh nghiệp", "Điểm giao dịch, đánh giá nông dân, xác thực pháp lý và huy hiệu đối tác.", "Đã có nền tảng dữ liệu/MVP"),
    ("Kết nối tài chính/chính sách", "Tận dụng dữ liệu rủi ro, giao dịch và lịch sử để hỗ trợ tiếp cận nguồn lực.", "Mục tiêu pilot/đối tác"),
    ("Hộ chiếu số trang trại", "Chuẩn hóa hồ sơ mùa vụ, vùng trồng và kết quả thực tế theo thời gian.", "Lộ trình sau pilot"),
], widths=[1.55, 3.65, 1.9], font_size=8.4)

doc.add_heading("1.3. Mức giá và mô hình thu", level=2)
add_p(doc, "Theo kế hoạch tài chính hiện tại, doanh nghiệp là bên trả phí để duy trì hạ tầng dùng chung, trong khi nông dân được ưu tiên tiếp cận các chức năng thiết yếu với chi phí thấp hoặc miễn phí trong giai đoạn mở rộng người dùng.")
bullet(doc, "Gói nền tảng doanh nghiệp mục tiêu: 12 triệu đồng/doanh nghiệp/tháng, cần được kiểm chứng bằng phỏng vấn willingness-to-pay và pilot có trả phí.")
bullet(doc, "Hoa hồng dịch vụ/giao dịch: kế hoạch đang mô hình hóa khoảng 5% với giao dịch trả trước và 6% với giao dịch trả sau; chỉ áp dụng khi có giá trị phát sinh và thỏa thuận minh bạch.")
bullet(doc, "Nguồn thu bảo hiểm nông nghiệp đang để mở trong dự báo chính; chỉ triển khai sau khi có đối tác được cấp phép và mô hình tuân thủ phù hợp.")

doc.add_heading("1.4. Điểm khác biệt cốt lõi", level=2)
number(doc, "Khép kín vòng lặp dữ liệu – hành động – kết quả, thay vì chỉ hiển thị cảnh báo.")
number(doc, "Kết nối ba nhóm có lợi ích bổ sung: người sản xuất, bên mua/cung ứng và tổ chức tài chính.")
number(doc, "Xây uy tín doanh nghiệp từ giao dịch đã xác thực, đánh giá sau giao dịch, GPS và hồ sơ pháp lý; giảm phụ thuộc vào quảng cáo hoặc lời truyền miệng.")
number(doc, "Thiết kế từ bối cảnh ĐBSCL, ưu tiên giao diện di động và quyết định ngắn gọn, phù hợp cách sử dụng thực tế của nông dân.")

doc.add_heading("1.5. Ý nghĩa tên gọi", level=2)
add_p(doc, "“CHAMP” gợi “Chăm” – sự chăm sóc bền bỉ; “MẠNH” đồng âm gợi “Mặn” – rủi ro khí hậu dự án lựa chọn giải quyết; “CHAT” gợi sự đối thoại và “chắt chiu” – tinh thần tối ưu từng nguồn lực. Tên gọi vừa có năng lượng chiến thắng, vừa gắn với đời sống nông nghiệp miền Tây.")

page_break(doc)

doc.add_heading("2. VẤN ĐỀ VÀ TÍNH CẤP THIẾT", level=1)
doc.add_heading("2.1. Một rủi ro khí hậu kéo theo nhiều quyết định kinh tế", level=2)
add_p(doc, "Xâm nhập mặn không chỉ là một con số môi trường. Sai thời điểm lấy nước có thể kéo theo thiệt hại cây trồng, tăng chi phí xử lý, gián đoạn kế hoạch thu mua và làm suy yếu khả năng trả nợ. Trong khi đó, thông tin thường nằm ở nhiều nguồn, khó chuyển thành quyết định phù hợp cho từng nông hộ.")
add_p(doc, "Theo số liệu được đội dẫn trong Application Form, mùa khô 2024–2025 (tính đến 05/03/2025), ranh mặn 4‰ đã xâm nhập khoảng 42–60 km ở các cửa sông Cửu Long và 52–57 km ở hệ thống Vàm Cỏ. Trong mùa khô 2023–2024, khoảng 56.260 ha lúa Đông Xuân và 43.300 ha cây ăn trái nằm trong vùng có nguy cơ ảnh hưởng. Các con số này cho thấy nhu cầu chuyển từ phản ứng bị động sang quản trị rủi ro có dữ liệu.")

doc.add_heading("2.2. Ba khoảng trống đang tồn tại", level=2)
add_table(doc, ["Nhóm", "Khoảng trống", "Hệ quả"], [
    ("Nông dân", "Thiếu thông tin hành động đúng lúc; đầu vào, vốn và đầu ra rời rạc.", "Tăng chi phí, ra quyết định muộn, dễ bị ép giá và khó chứng minh năng lực sản xuất."),
    ("Doanh nghiệp", "Thiếu dữ liệu minh bạch về vùng nguyên liệu, thời điểm, nhu cầu và độ tin cậy của giao dịch.", "Tăng chi phí tìm nguồn, khó lập kế hoạch thu mua, chất lượng không đồng đều."),
    ("Tổ chức tài chính/chính sách", "Thiếu dữ liệu vận hành liên tục để nhìn rủi ro và kết quả sử dụng vốn.", "Thẩm định chậm, sản phẩm khó sát mùa vụ, chi phí giám sát cao."),
], widths=[1.2, 3.0, 2.9], font_size=8.7)

doc.add_heading("2.3. Câu hỏi thiết kế", level=2)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Làm thế nào để một cảnh báo độ mặn trở thành một chuỗi hành động có người chịu trách nhiệm, có nguồn lực đi kèm và có kết quả được ghi nhận?")
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = RGBColor.from_string(GREEN)

doc.add_heading("2.4. Cách CHAMP MẠNH CHAT giải quyết", level=2)
add_table(doc, ["Bước", "Hệ thống thực hiện", "Kết quả cần đo"], [
    ("1. Nhận biết", "Tổng hợp dữ liệu dự báo, hiển thị bản đồ và cảnh báo theo khu vực.", "Tỷ lệ cảnh báo được đọc; độ phủ địa bàn."),
    ("2. Quyết định", "Chuyển mức rủi ro thành khuyến nghị ngắn gọn, gắn với thời điểm và cây trồng.", "Tỷ lệ người dùng hiểu và thực hiện khuyến nghị."),
    ("3. Kết nối", "Kết nối nhu cầu đầu vào, thu mua và nguồn lực tài chính phù hợp.", "Số kết nối đủ điều kiện; thời gian từ nhu cầu đến phản hồi."),
    ("4. Xác thực", "Ghi nhận GPS, trạng thái giao dịch, khối lượng, đánh giá và hồ sơ pháp lý.", "Tỷ lệ giao dịch hoàn tất; sai lệch cam kết; khiếu nại."),
    ("5. Học hỏi", "Dùng dữ liệu kết quả để cải thiện khuyến nghị, uy tín và phân bổ nguồn lực.", "Độ chính xác dự báo và chất lượng quyết định tăng qua từng vụ."),
], widths=[0.8, 4.2, 2.1], font_size=8.5)

page_break(doc)

# Part 2
doc.add_heading("PHẦN 2 · LIÊN KẾT VỚI CHỦ ĐỀ CUỘC THI", level=1)
doc.add_heading("1. Tối ưu hóa nguồn lực – giảm lãng phí – tối đa hóa giá trị", level=2)
add_table(doc, ["Nguồn lực", "Lãng phí hiện nay", "Cơ chế tối ưu của dự án"], [
    ("Nước ngọt", "Lấy nước sai thời điểm hoặc phản ứng quá muộn trước xâm nhập mặn.", "Cảnh báo theo địa bàn, hỗ trợ lựa chọn thời điểm và phương án ứng phó."),
    ("Vật tư & lao động", "Đầu tư không khớp rủi ro/mùa vụ; mua manh mún; khó điều phối.", "Gắn nhu cầu vào hồ sơ mùa vụ và kết nối nhà cung ứng phù hợp."),
    ("Nông sản", "Thu hoạch/thu mua thiếu phối hợp, giảm phẩm cấp, tăng thất thoát sau thu hoạch.", "Doanh nghiệp đăng nhu cầu rõ ràng; nông dân đăng ký; lịch và kết quả được ghi nhận."),
    ("Vốn", "Nguồn lực đến chậm hoặc phân bổ thiếu căn cứ dữ liệu.", "Tạo hồ sơ dữ liệu về rủi ro, nhu cầu, giao dịch và kết quả để hỗ trợ thẩm định."),
    ("Thông tin & niềm tin", "Dữ liệu phân mảnh; uy tín phụ thuộc truyền miệng.", "Một hồ sơ giao dịch có GPS, đánh giá sau giao dịch và xác thực pháp lý."),
], widths=[1.2, 2.7, 3.2], font_size=8.6)

doc.add_heading("2. Đóng góp cho các Mục tiêu Phát triển Bền vững", level=2)
add_table(doc, ["SDG", "Đóng góp trực tiếp của CHAMP MẠNH CHAT"], [
    ("SDG 2 – Xóa đói", "Tăng khả năng chống chịu của sản xuất nông nghiệp, cải thiện tiếp cận thị trường và nguồn lực."),
    ("SDG 8 – Việc làm & tăng trưởng", "Giảm chi phí kết nối, mở rộng cơ hội kinh doanh địa phương và tạo việc làm số trong nông nghiệp."),
    ("SDG 9 – Công nghiệp, đổi mới & hạ tầng", "Xây hạ tầng dữ liệu và quyết định số cho chuỗi giá trị nông nghiệp."),
    ("SDG 12 – Sản xuất, tiêu dùng có trách nhiệm", "Điều phối nhu cầu – sản lượng – thu mua để giảm đầu vào lãng phí và thất thoát."),
    ("SDG 13 – Hành động khí hậu", "Chuyển dữ liệu rủi ro khí hậu thành hành động thích ứng có thể đo lường."),
], widths=[2.15, 4.95], font_size=8.8)
add_p(doc, "Dự án có liên hệ bổ trợ với SDG 3 và SDG 4 thông qua thực phẩm an toàn hơn, nội dung hướng dẫn dễ tiếp cận và nâng cao năng lực số; tuy nhiên hồ sơ ưu tiên các SDG có quan hệ trực tiếp để tránh dàn trải tác động.")

doc.add_heading("3. Công nghệ bền vững", level=2)
bullet(doc, "Dữ liệu trước thiết bị: ưu tiên tận dụng nguồn dữ liệu hiện có, bản đồ mở và kiến trúc cloud linh hoạt trước khi đầu tư IoT diện rộng.")
bullet(doc, "Thử nghiệm nhỏ trước mở rộng: chỉ mở thêm địa bàn/cây trồng khi bộ chỉ số pilot chứng minh hiệu quả.")
bullet(doc, "Quyền riêng tư theo mục đích: GPS chi tiết chỉ được lưu khi người dùng chủ động; phân quyền truy cập và nhật ký dữ liệu phải được kiểm soát.")
bullet(doc, "Đo tác động theo vụ: không đánh giá bằng lượt tải đơn thuần, mà bằng quyết định đúng lúc, giao dịch hoàn tất, tài nguyên tiết kiệm và thu nhập được bảo vệ.")

doc.add_heading("4. Mô hình kinh tế xanh", level=2)
add_p(doc, "CHAMP MẠNH CHAT theo đuổi mô hình “dữ liệu dùng chung – giá trị cùng tạo”: dữ liệu không được bán như một tài sản thô của nông dân. Giá trị kinh tế đến từ năng lực ra quyết định, kết nối và xác thực tốt hơn. Doanh nghiệp trả phí vì giảm chi phí tìm nguồn và tăng chất lượng cung ứng; đối tác tài chính nhận lợi ích từ dữ liệu vận hành đáng tin; nông dân được hưởng dịch vụ thiết yếu và cơ hội thị trường.")

page_break(doc)

# Part 3
doc.add_heading("PHẦN 3 · PHÂN TÍCH THỊ TRƯỜNG", level=1)
doc.add_heading("1. Tổng quan và phân khúc", level=2)
add_p(doc, "Thị trường mục tiêu không được định nghĩa bằng “toàn bộ nông nghiệp Việt Nam” ngay từ đầu. Dự án lựa chọn một thị trường có thể phục vụ và kiểm chứng: chuỗi giá trị cây ăn trái giá trị cao tại ĐBSCL, bắt đầu bằng mô hình sầu riêng ở Cần Thơ và vùng lân cận, nơi rủi ro nước – độ mặn – chất lượng – thời điểm thu mua liên kết chặt với giá trị kinh tế.")
add_table(doc, ["Phân khúc", "Nhu cầu cấp thiết", "Giá trị chào bán", "Ưu tiên"], [
    ("Nông hộ/nhà vườn", "Cảnh báo dễ hiểu, đầu ra tin cậy, tiếp cận vốn/đầu vào.", "Ứng dụng miễn phí hoặc trợ giá; hồ sơ và kết nối giao dịch.", "Người dùng lõi"),
    ("Doanh nghiệp/vựa thu mua", "Nguồn cung đúng vùng, đúng thời điểm, giảm chi phí chạy địa bàn.", "Đăng nhu cầu, quản lý lead/giao dịch, uy tín xác thực, dữ liệu tổng hợp.", "Khách hàng trả phí đầu tiên"),
    ("Hợp tác xã", "Điều phối thành viên, tiêu chuẩn hóa hồ sơ và đầu ra.", "Bảng điều phối nhóm, dữ liệu vùng, kết nối doanh nghiệp.", "Đối tác triển khai"),
    ("Ngân hàng/chính sách", "Dữ liệu rủi ro, mục đích vốn và kết quả sử dụng.", "API/báo cáo đồng thuận, hỗ trợ sàng lọc và giám sát.", "Đối tác giai đoạn 2"),
    ("Nhà cung ứng/bảo hiểm", "Tiếp cận đúng nhu cầu và đúng thời điểm.", "Kênh phân phối có dữ liệu và đo lường chuyển đổi.", "Mở rộng có kiểm soát"),
], widths=[1.4, 2.3, 2.6, 0.8], font_size=7.9)

doc.add_heading("2. Quy mô thị trường theo phương pháp bottom-up", level=2)
add_p(doc, "Để tránh ước lượng TAM quá rộng, đội dùng số doanh nghiệp có thể ký hợp đồng làm đơn vị doanh thu. Mô hình kế hoạch đặt mục tiêu 10 doanh nghiệp trả phí ở Năm 1, 15 ở Năm 2, 25 ở Năm 3 và 100 ở Năm 7; mức phí mục tiêu 12 triệu đồng/tháng. Đây là mục tiêu thương mại cần được kiểm chứng, không phải dữ liệu thị trường đã xác nhận.")
add_table(doc, ["Tầng thị trường", "Định nghĩa thực dụng", "Cách xác thực"], [
    ("Beachhead", "Doanh nghiệp/vựa thu mua sầu riêng tại Cần Thơ và cụm lân cận có nhu cầu nguồn cung lặp lại.", "20–30 phỏng vấn; 3–5 MoU; tối thiểu 2 pilot có giao dịch thật."),
    ("SAM gần hạn", "Doanh nghiệp thu mua và HTX cây ăn trái ở 13 tỉnh/thành ĐBSCL.", "Mở rộng theo cụm chuỗi giá trị sau khi đạt KPI pilot."),
    ("TAM dài hạn", "Các chuỗi nông sản nhạy cảm khí hậu tại Việt Nam và thị trường tương đồng.", "Chuẩn hóa dữ liệu, quy trình và đối tác; không mở rộng trước khi unit economics ổn định."),
], widths=[1.2, 3.7, 2.2], font_size=8.5)

doc.add_heading("3. Đối thủ và giải pháp thay thế", level=2)
add_table(doc, ["Nhóm giải pháp", "Điểm mạnh", "Khoảng trống", "Khác biệt của dự án"], [
    ("Ứng dụng thời tiết/quan trắc", "Dữ liệu nhanh, phạm vi rộng.", "Thường dừng ở hiển thị; ít gắn với giao dịch và nguồn lực.", "Chuyển cảnh báo thành luồng hành động và kết quả."),
    ("Sàn/chợ nông sản số", "Kết nối cung – cầu.", "Thiếu ngữ cảnh rủi ro khí hậu và xác thực sau giao dịch.", "Gắn nhu cầu thu mua với địa bàn, GPS và uy tín."),
    ("Nhóm Zalo/Facebook/môi giới", "Quen thuộc, tốc độ lan truyền cao.", "Dữ liệu rời rạc, khó truy vết, rủi ro thông tin và ép giá.", "Quy trình có cấu trúc, hồ sơ và lịch sử giao dịch."),
    ("Phần mềm quản lý nông trại", "Quản trị chuyên sâu cho đơn vị lớn.", "Chi phí/độ phức tạp cao với nông hộ; thường không có tài chính/thu mua.", "Thiết kế đơn giản, ba bên, triển khai theo cụm."),
    ("Ứng dụng tín dụng/bảo hiểm", "Chuyên môn tài chính, quy trình tuân thủ.", "Thiếu dữ liệu mùa vụ và giao dịch thực địa liên tục.", "Đóng vai trò lớp dữ liệu/kết nối; không thay thế tổ chức được cấp phép."),
], widths=[1.25, 1.65, 2.2, 2.0], font_size=7.8)

page_break(doc)

doc.add_heading("4. Khách hàng mục tiêu và chân dung", level=1)
doc.add_heading("4.1. Khách hàng trả phí ưu tiên: doanh nghiệp/vựa thu mua uy tín", level=2)
add_table(doc, ["Tiêu chí", "Chân dung"], [
    ("Địa lý", "Cần Thơ và các tỉnh ĐBSCL có vùng cây ăn trái tập trung; mở rộng theo cụm cung ứng."),
    ("Quy mô", "Doanh nghiệp/vựa có hoạt động thu mua lặp lại, đội chạy địa bàn và nhu cầu chuẩn hóa nguồn cung."),
    ("Tâm lý", "Muốn tách mình khỏi thương lái thiếu uy tín; sẵn sàng minh bạch để xây thương hiệu."),
    ("Hành vi mua", "Chỉ trả phí khi có lead phù hợp, giảm chi phí tìm nguồn hoặc cải thiện tỷ lệ hoàn tất giao dịch."),
    ("Job-to-be-done", "“Giúp tôi tìm đúng nhà vườn, đúng sản lượng, đúng thời điểm và chứng minh rằng tôi là người mua đáng tin.”"),
], widths=[1.35, 5.75], font_size=9)

doc.add_heading("4.2. Người dùng lõi: nông hộ/nhà vườn", level=2)
add_table(doc, ["Tiêu chí", "Chân dung"], [
    ("Địa lý", "ĐBSCL; ưu tiên nơi có rủi ro độ mặn và vùng cây trồng giá trị cao."),
    ("Thiết bị", "Dùng điện thoại thông minh; ưu tiên giao diện ít bước, chữ rõ, có thể đọc nội dung bằng giọng nói."),
    ("Nỗi đau", "Thông tin chậm/khó hiểu, đầu ra thiếu chắc chắn, khó đánh giá người mua, thiếu hồ sơ để tiếp cận nguồn lực."),
    ("Động lực", "Bảo vệ mùa vụ, bán được đúng thời điểm, giao dịch minh bạch và tăng uy tín qua lịch sử thực tế."),
    ("Rào cản", "Niềm tin, kỹ năng số, lo ngại chia sẻ vị trí/dữ liệu, thói quen giao dịch ngoài nền tảng."),
], widths=[1.35, 5.75], font_size=9)

doc.add_heading("4.3. Đơn vị ra quyết định và người ảnh hưởng", level=2)
add_p(doc, "Trong pilot, người sử dụng không đồng nghĩa người ký hợp đồng. Doanh nghiệp là người trả phí; HTX/chính quyền địa phương/cán bộ khuyến nông là bên tạo niềm tin và hỗ trợ onboarding; nông dân là người tạo dữ liệu và giá trị mạng lưới; tổ chức tài chính là đối tác mua/đồng tài trợ giải pháp ở giai đoạn sau.")

doc.add_heading("5. Chiến lược tiếp cận thị trường", level=2)
add_table(doc, ["Giai đoạn", "Cách làm", "Chỉ số quyết định"], [
    ("0–3 tháng", "Phỏng vấn nhà vườn, vựa, HTX; kiểm thử nhu cầu và thông điệp; hoàn thiện bảo mật/dữ liệu.", "≥30 phỏng vấn; ≥3 đối tác pilot; danh sách 100 nông hộ tiềm năng."),
    ("4–6 tháng", "Pilot 1 cụm địa bàn, 1 cây trồng; onboarding trực tiếp qua HTX/đối tác.", "≥50 nông hộ hoạt động; ≥30 đăng ký thu mua; ≥10 giao dịch hoàn tất."),
    ("7–12 tháng", "Đo hiệu quả, điều chỉnh giá, chuẩn hóa playbook; thử gói trả phí doanh nghiệp.", "≥2 doanh nghiệp trả phí; giữ chân 60% sau 3 tháng; tỷ lệ hoàn tất ≥70%."),
    ("Năm 2", "Mở rộng theo cụm ĐBSCL; thêm đối tác dữ liệu/tài chính khi có bằng chứng.", "15 doanh nghiệp theo kế hoạch; dữ liệu đủ cho báo cáo rủi ro/chuỗi cung ứng."),
], widths=[1.05, 4.45, 1.6], font_size=8.3)

doc.add_heading("6. Kênh truyền thông", level=2)
bullet(doc, "Kênh tin cậy: HTX, tổ hợp tác, cán bộ khuyến nông, hội nông dân và doanh nghiệp đầu chuỗi.")
bullet(doc, "Nội dung theo mùa vụ: bản tin độ mặn, checklist hành động, câu chuyện giao dịch minh bạch, video ngắn bằng ngôn ngữ địa phương.")
bullet(doc, "Chứng minh thay vì quảng cáo: công bố dashboard pilot gồm số cảnh báo hữu ích, giao dịch hoàn tất, thời gian tìm nguồn và phản hồi người dùng.")

page_break(doc)

# Part 4
doc.add_heading("PHẦN 4 · NĂNG LỰC NỘI BỘ VÀ SWOT", level=1)
add_table(doc, ["ĐIỂM MẠNH", "ĐIỂM YẾU"], [
    ("• MVP có nhiều luồng cốt lõi thay vì chỉ là slide ý tưởng.\n• Bài toán ba bên tạo dữ liệu và hiệu ứng mạng lưới.\n• Đội hiểu bối cảnh ĐBSCL và có tinh thần xây sản phẩm.\n• Kiến trúc Supabase/React giúp thử nghiệm nhanh, chi phí ban đầu linh hoạt.",
     "• Chưa có doanh thu và chưa có đối tác chính thức.\n• Mô hình tài chính phụ thuộc mạnh vào số DN trả phí.\n• Đội sinh viên còn thiếu năng lực bán hàng B2B, pháp lý dữ liệu và vận hành nông nghiệp thực địa.\n• Một số mô-đun tài chính/IDSS còn ở lộ trình."),
], widths=[3.55, 3.55], font_size=8.8)
add_table(doc, ["CƠ HỘI", "THÁCH THỨC"], [
    ("• Nhu cầu thích ứng khí hậu và truy xuất chuỗi cung ứng tăng.\n• Doanh nghiệp uy tín cần công cụ giảm chi phí tìm vùng nguyên liệu.\n• Ngân hàng/đơn vị chính sách cần dữ liệu thực tế để thiết kế sản phẩm phù hợp.\n• Có thể nhân rộng sang hạn, lũ, sâu bệnh sau khi chuẩn hóa mô hình.",
     "• Dữ liệu dự báo có sai số; kết nối Internet và kỹ năng số không đồng đều.\n• Người dùng có thể đưa giao dịch ra ngoài nền tảng.\n• Rủi ro quyền riêng tư GPS, dữ liệu cá nhân và quy định đối với tài chính/bảo hiểm.\n• Cần đạt mật độ cung–cầu đủ lớn ở từng địa bàn."),
], widths=[3.55, 3.55], font_size=8.8)

doc.add_heading("Chiến lược từ SWOT", level=2)
add_table(doc, ["Chiến lược", "Hành động"], [
    ("SO – Dùng MVP để mở cửa pilot", "Demo luồng end-to-end bằng dữ liệu thật; ưu tiên đối tác có nhu cầu thu mua lặp lại."),
    ("WO – Bù năng lực còn thiếu", "Mời cố vấn nông nghiệp, pháp lý dữ liệu, tài chính; hợp tác HTX để vận hành thực địa."),
    ("ST – Chống rủi ro niềm tin", "Minh bạch nguồn dữ liệu, khoảng tin cậy; chỉ mở đánh giá sau giao dịch; quy trình xử lý khiếu nại."),
    ("WT – Giới hạn phạm vi", "Một địa bàn – một cây trồng – một bộ KPI; dừng hoặc điều chỉnh nếu không đạt cổng kiểm chứng."),
], widths=[2.05, 5.05], font_size=8.7)

doc.add_heading("Ma trận rủi ro trọng yếu", level=2)
add_table(doc, ["Rủi ro", "Mức", "Giảm thiểu", "Tín hiệu cảnh báo sớm"], [
    ("Dự báo sai/không đủ dữ liệu", "Cao", "Hiển thị khoảng tin cậy; đối chiếu trạm; khuyến nghị không thay thế chuyên gia.", "Sai lệch tăng; người dùng bỏ qua cảnh báo."),
    ("Không đủ thanh khoản cung–cầu", "Cao", "Onboard theo cụm; ký doanh nghiệp neo; chỉ mở địa bàn khi đủ mật độ.", "Lead không phản hồi; giao dịch hoàn tất thấp."),
    ("Không chấp nhận mức phí", "Cao", "Pilot có trả phí nhỏ; gói theo quy mô; định giá dựa trên chi phí tiết kiệm.", "Nhiều dùng thử, ít chuyển đổi trả phí."),
    ("Rò rỉ/lạm dụng dữ liệu", "Cao", "Consent theo mục đích; RLS; tối thiểu hóa dữ liệu; nhật ký truy cập; xóa dữ liệu.", "Truy cập bất thường; khiếu nại quyền riêng tư."),
    ("Giao dịch ra ngoài nền tảng", "TB", "Uy tín, bảo chứng, quy trình và ưu đãi chỉ có khi giao dịch được ghi nhận.", "Nhiều liên hệ nhưng ít hoàn tất trong app."),
    ("Phụ thuộc đối tác tài chính", "TB", "Tách mô-đun; không hứa phê duyệt vốn; duy trì nhiều đối tác tiềm năng.", "Thời gian tích hợp kéo dài."),
], widths=[1.35, 0.55, 3.35, 1.85], font_size=7.7)

page_break(doc)

# Part 5
doc.add_heading("PHẦN 5 · DỰ BÁO TÀI CHÍNH", level=1)
doc.add_heading("1. Cơ sở lập kế hoạch", level=2)
add_table(doc, ["Giả định", "Giá trị kế hoạch", "Cách kiểm chứng"], [
    ("Tổng đầu tư ban đầu", "3,010 tỷ đồng", "Báo giá phát triển, kế hoạch nhân sự, chi phí pilot theo giai đoạn."),
    ("Cấu trúc vốn", "60% vốn chủ; 40% vốn vay", "Chỉ thực hiện khi có nguồn vốn phù hợp và khả năng trả nợ."),
    ("Chi phí vốn bình quân", "WACC 14,32%", "Cập nhật theo điều kiện vốn thực tế trước khi huy động."),
    ("Phí nền tảng DN", "12 triệu đồng/DN/tháng", "Pilot trả phí và phỏng vấn willingness-to-pay."),
    ("DN đối tác", "10 ở Năm 1 → 100 ở Năm 7", "Pipeline có tên, tỷ lệ chuyển đổi và cohort giữ chân."),
    ("Giá trị giao dịch", "Tăng theo số DN và tỷ lệ GMV đi qua nền tảng", "Chỉ ghi nhận giao dịch xác thực, không tính lead chưa hoàn tất."),
], widths=[2.0, 1.65, 3.45], font_size=8.5)

doc.add_heading("2. Kế hoạch doanh thu – chi phí – lợi nhuận", level=2)
add_p(doc, "Đơn vị: triệu đồng. Số liệu được tổng hợp trực tiếp từ báo cáo tài chính đội cung cấp; làm tròn để trình bày.")
add_table(doc, ["Chỉ tiêu", "Năm 1", "Năm 2", "Năm 3", "Năm 4"], [
    ("Doanh nghiệp đối tác", "10", "15", "25", "40"),
    ("Tổng doanh thu", "1.580", "2.429", "4.188", "6.992"),
    ("Chi phí vận hành", "1.934", "2.389", "3.381", "5.154"),
    ("Lợi nhuận ròng", "(1.039)", "(614)", "142", "1.147"),
    ("Biên lợi nhuận ròng", "-65,8%", "-25,3%", "3,4%", "16,4%"),
], widths=[2.2, 1.2, 1.2, 1.2, 1.2], font_size=8.7)
doc.add_paragraph()
add_table(doc, ["Chỉ tiêu", "Năm 5", "Năm 6", "Năm 7"], [
    ("Doanh nghiệp đối tác", "60", "80", "100"),
    ("Tổng doanh thu", "12.091", "16.368", "20.691"),
    ("Chi phí vận hành", "7.466", "9.538", "11.459"),
    ("Lợi nhuận ròng", "3.400", "5.258", "7.180"),
    ("Biên lợi nhuận ròng", "28,1%", "32,1%", "34,7%"),
], widths=[2.65, 1.45, 1.45, 1.45], font_size=8.9)

doc.add_heading("3. Hiệu quả vốn trong kịch bản cơ sở", level=2)
add_table(doc, ["Chỉ tiêu", "Kết quả", "Cách đọc thận trọng"], [
    ("NPV", "4,442 tỷ đồng", "Dương ở WACC 14,32%; phụ thuộc vào tốc độ tăng doanh nghiệp và giao dịch."),
    ("IRR", "36,11%", "Cao hơn WACC trong mô hình; cần kiểm tra lại sau dữ liệu pilot."),
    ("MIRR", "30,85%", "Giả định tái đầu tư ở WACC, thận trọng hơn IRR."),
    ("Hoàn vốn", "Dòng tiền chủ sở hữu chưa chiết khấu chuyển dương trong Năm 5; chiết khấu chuyển dương trong Năm 6.", "Phù hợp với dự án hạ tầng số B2B nhưng đòi hỏi kỷ luật vốn 24 tháng đầu."),
], widths=[1.4, 2.2, 3.5], font_size=8.4)

doc.add_heading("4. Kỷ luật tài chính và cổng kiểm chứng", level=2)
bullet(doc, "Không giải ngân toàn bộ 1,8 tỷ đồng phát triển ứng dụng ngay từ đầu; chia theo mốc sản phẩm và hợp đồng pilot.")
bullet(doc, "Trước khi tăng nhân sự, chứng minh ba chỉ số: doanh nghiệp trả phí, tỷ lệ hoàn tất giao dịch và giữ chân cohort.")
bullet(doc, "Phí 12 triệu đồng/tháng chỉ trở thành giá chuẩn khi khách hàng xác nhận giá trị tiết kiệm lớn hơn phí; nếu không, chuyển sang gói theo quy mô/usage.")
bullet(doc, "Không ghi nhận doanh thu bảo hiểm trong kịch bản cơ sở khi chưa có đối tác và mô hình tuân thủ.")
bullet(doc, "Báo cáo theo cả chỉ số tăng trưởng và chỉ số tác động: doanh thu không được tăng bằng cách đánh đổi quyền dữ liệu hay chất lượng khuyến nghị.")

page_break(doc)
doc.add_heading("5. Nhu cầu nguồn lực đề xuất", level=2)
add_table(doc, ["Hạng mục", "Tỷ trọng định hướng", "Đầu ra bắt buộc"], [
    ("Sản phẩm & dữ liệu", "40%", "Ổn định MVP, dữ liệu cảnh báo, bảo mật, dashboard pilot."),
    ("Pilot & vận hành địa bàn", "25%", "Onboarding, hỗ trợ người dùng, đo kết quả mùa vụ."),
    ("Phát triển thị trường", "15%", "Pipeline doanh nghiệp, tài liệu bán hàng, hợp đồng thử nghiệm."),
    ("Pháp lý & quản trị dữ liệu", "10%", "Consent, điều khoản sử dụng, đánh giá bảo mật và quy trình khiếu nại."),
    ("Dự phòng", "10%", "Ứng phó trễ dữ liệu, thay đổi mùa vụ và chi phí tích hợp."),
], widths=[2.0, 1.35, 3.75], font_size=8.7)

# Team and roadmap
doc.add_heading("PHẦN 6 · ĐỘI NGŨ VÀ KẾ HOẠCH THỰC THI", level=1)
doc.add_heading("1. Đội ngũ", level=2)
add_table(doc, ["Thành viên", "Vai trò", "Thông tin"], [
    ("Phan Khắc Anh Tuấn", "Trưởng nhóm", "UEH · MSSV 31241024275 · 0585 708 372"),
    ("Lương Thị Bạch Dương", "Thành viên", "UEH · MSSV 31231024482"),
    ("Dịp Minh Châu", "Thành viên", "UEH · MSSV 31241024259"),
    ("Nguyễn Ngọc Đường Nghi", "Thành viên", "UEH · MSSV 31241021416"),
    ("Phạm Đình Quang", "Thành viên", "UEH · MSSV 31241024277"),
], widths=[2.3, 1.25, 3.55], font_size=8.8)
add_p(doc, "Đội đã đồng hành khoảng 1 năm 3 tháng theo Application Form. Điểm cần bổ sung ngay không phải thêm số lượng thành viên, mà là ba năng lực cố vấn: nông học/thủy văn ĐBSCL, bán hàng B2B nông nghiệp và pháp lý dữ liệu – tài chính.")

doc.add_heading("2. Phân công theo kết quả", level=2)
add_table(doc, ["Trục công việc", "Trách nhiệm đầu ra", "Cơ chế kiểm soát"], [
    ("Sản phẩm & công nghệ", "MVP ổn định, dữ liệu đúng, trải nghiệm ít bước, bảo mật.", "Sprint 2 tuần; checklist release; log lỗi và phản hồi người dùng."),
    ("Thị trường & đối tác", "Phỏng vấn, MoU, pilot và hợp đồng trả phí.", "CRM pipeline; review tỷ lệ chuyển đổi hàng tuần."),
    ("Vận hành & cộng đồng", "Onboarding nông hộ, hỗ trợ giao dịch, thu thập kết quả.", "SLA hỗ trợ; dashboard cohort; nhật ký sự cố."),
    ("Tài chính & đo tác động", "Ngân sách, unit economics, KPI môi trường/xã hội.", "Báo cáo tháng; cổng go/no-go theo giai đoạn."),
], widths=[1.55, 3.8, 1.75], font_size=8.4)

doc.add_heading("3. Lộ trình 18 tháng", level=2)
add_table(doc, ["Thời gian", "Mục tiêu", "Sản phẩm bàn giao", "Cổng quyết định"], [
    ("0–3 tháng", "Chuẩn bị pilot", "MVP ổn định; quy trình dữ liệu; 3–5 đối tác; thiết kế đo lường.", "Có địa bàn, cây trồng, đối tác neo và người chịu trách nhiệm."),
    ("4–6 tháng", "Pilot thực địa", "Cảnh báo + thu mua + đánh giá end-to-end; dashboard pilot.", "≥10 giao dịch hoàn tất; người dùng xác nhận giá trị."),
    ("7–12 tháng", "Thử trả phí", "Gói doanh nghiệp; playbook onboarding; báo cáo tác động.", "≥2 DN trả phí; retention ≥60%; tỷ lệ hoàn tất ≥70%."),
    ("13–18 tháng", "Chuẩn hóa & mở rộng", "Mở thêm cụm địa bàn; dữ liệu cho đối tác tài chính; hộ chiếu số bản thử nghiệm.", "Unit economics dương ở cohort trưởng thành; kiểm soát dữ liệu đạt yêu cầu."),
], widths=[1.0, 1.4, 3.2, 1.5], font_size=8.1)

doc.add_heading("4. Bộ chỉ số thành công", level=2)
add_table(doc, ["Nhóm KPI", "Chỉ số"], [
    ("Sản phẩm", "Tỷ lệ cảnh báo được mở; thời gian đến hành động; tỷ lệ người dùng hoàn thành luồng chính."),
    ("Thị trường", "Doanh nghiệp trả phí; CAC; retention; số lead đủ điều kiện; tỷ lệ giao dịch hoàn tất."),
    ("Niềm tin", "Tỷ lệ giao dịch có GPS/đánh giá; khiếu nại; thời gian xử lý; tỷ lệ doanh nghiệp xác thực."),
    ("Tác động", "Diện tích/nông hộ được hỗ trợ; quyết định lấy nước đúng lúc; đầu vào hoặc tổn thất giảm; thu nhập được bảo vệ."),
    ("Tài chính", "MRR, gross margin, burn rate, runway, NPV cập nhật và sai lệch so với kế hoạch."),
], widths=[1.35, 5.75], font_size=8.8)

page_break(doc)

# Q&A required
doc.add_heading("PHẦN 7 · TRẢ LỜI CÂU HỎI CỦA BAN TỔ CHỨC", level=1)
doc.add_heading("1. Động lực nào khiến đội phát triển dự án? Nguồn lực nào đang bị lãng phí?", level=2)
add_p(doc, "Đội bắt đầu từ quan sát rằng người nông dân không thiếu thông tin, nhưng thiếu một con đường rõ ràng từ thông tin đến quyết định và nguồn lực. Nước ngọt, vật tư, thời gian lao động, nông sản, vốn và cả niềm tin đang bị lãng phí vì dữ liệu đến muộn, rời rạc hoặc không gắn với người có khả năng hành động. CHAMP MẠNH CHAT kết nối dữ liệu khí hậu với nhu cầu thật, giao dịch thật và kết quả thật để mỗi nguồn lực được dùng đúng nơi, đúng lúc.")

doc.add_heading("2. Đội mong muốn học hỏi và đạt được gì tại cuộc thi?", level=2)
add_p(doc, "Đội mong kiểm chứng ba câu hỏi sống còn: nông dân có thực sự thay đổi hành động khi nhận khuyến nghị không; doanh nghiệp có trả phí để giảm chi phí tìm nguồn và tăng tỷ lệ giao dịch thành công không; dữ liệu tạo ra có đủ tin cậy để một đối tác tài chính sử dụng không. Startup Zone là môi trường để nhận phản biện, chuẩn hóa mô hình, kết nối cố vấn/đối tác và chuyển MVP thành pilot có bằng chứng.")

doc.add_heading("3. Ai là khách hàng chính và vì sao họ lựa chọn dự án?", level=2)
add_p(doc, "Khách hàng trả phí đầu tiên là doanh nghiệp/vựa thu mua uy tín tại ĐBSCL. Họ lựa chọn CHAMP MẠNH CHAT nếu nền tảng giúp giảm chi phí chạy địa bàn, rút ngắn thời gian tìm đúng nguồn cung, nâng tỷ lệ hoàn tất và tạo hồ sơ uy tín có thể chứng minh. Nông dân là người dùng lõi vì họ nhận cảnh báo, cơ hội đầu ra và cơ chế đánh giá người mua; chính giá trị hai phía tạo ra lợi thế mạng lưới.")

doc.add_heading("4. Kế hoạch truyền thông về tối ưu nguồn lực và bền vững", level=2)
add_p(doc, "Dự án truyền thông bằng bằng chứng theo mùa vụ: bản đồ rủi ro dễ hiểu, checklist hành động, câu chuyện giao dịch minh bạch và dashboard pilot. Thông điệp trung tâm là “đúng dữ liệu – đúng quyết định – đúng nguồn lực”. Kênh chính gồm HTX/hội nông dân/cán bộ khuyến nông để tạo tin cậy; video ngắn và cộng đồng trong ứng dụng để duy trì tương tác; báo cáo tác động cho doanh nghiệp và đối tác.")

doc.add_heading("5. Rào cản và phương án khắc phục", level=2)
add_p(doc, "Rào cản lớn nhất là niềm tin và mật độ mạng lưới, không chỉ là công nghệ. Dự án khắc phục bằng pilot theo cụm, có doanh nghiệp neo và đối tác địa phương; chỉ mở rộng khi đủ mật độ cung–cầu. Sai số dữ liệu được quản trị bằng nguồn dữ liệu minh bạch và khoảng tin cậy. Quyền riêng tư được bảo vệ bằng sự đồng thuận theo mục đích, phân quyền và tối thiểu hóa dữ liệu. Rủi ro pháp lý tài chính/bảo hiểm được xử lý bằng vai trò kết nối dữ liệu, không tự nhận chức năng của tổ chức được cấp phép.")

doc.add_heading("6. Điểm khác biệt so với mô hình tương tự", level=2)
add_p(doc, "Các giải pháp hiện hữu thường tối ưu một mắt xích: cảnh báo, chợ nông sản, quản lý nông trại hoặc tín dụng. CHAMP MẠNH CHAT thiết kế một vòng lặp khép kín: dữ liệu khí hậu tạo quyết định; quyết định kích hoạt kết nối nguồn lực; giao dịch tạo kết quả; kết quả quay lại làm giàu dữ liệu và uy tín. Điểm khác biệt chỉ có giá trị khi pilot chứng minh được ba kết quả: quyết định tốt hơn, giao dịch đáng tin hơn và chi phí kết nối thấp hơn.")

doc.add_heading("7. Vì sao doanh nghiệp chấp nhận được nông dân đánh giá?", level=2)
add_p(doc, "Doanh nghiệp làm ăn uy tín đang chịu chi phí lớn để tự chứng minh mình khác với người mua bẻ cọc hoặc ép giá. Đánh giá chỉ mở sau giao dịch xác thực, kết hợp tiêu chí giá, đúng hẹn/kỹ thuật và minh bạch/thái độ. Điểm uy tín cao giúp doanh nghiệp được ưu tiên hiển thị và được nhà vườn chất lượng chủ động tìm đến; đó là một tài sản thương hiệu đo được, đổi lại cho sự minh bạch.")

page_break(doc)

# Appendix
doc.add_heading("PHỤ LỤC · LOGIC SẢN PHẨM VÀ NGUYÊN TẮC THỰC THI", level=1)
doc.add_heading("A. Kiến trúc giá trị mục tiêu", level=2)
add_table(doc, ["Lớp", "Dữ liệu vào", "Xử lý", "Đầu ra"], [
    ("Nhận biết", "Dữ liệu độ mặn, địa bàn, thời gian", "Dự báo, lọc, trực quan hóa", "Cảnh báo và mức rủi ro"),
    ("Quyết định", "Cây trồng, mùa vụ, ngưỡng", "Quy tắc/khuyến nghị IDSS", "Việc cần làm và thời điểm"),
    ("Nguồn lực", "Nhu cầu vật tư, đầu ra, vốn", "Ghép nối điều kiện", "Đối tác/đề nghị phù hợp"),
    ("Xác thực", "GPS, trạng thái, khối lượng, đánh giá", "Kiểm tra quyền và tính hợp lệ", "Lịch sử giao dịch/điểm uy tín"),
    ("Học hỏi", "Kết quả thực tế và phản hồi", "Phân tích cohort/tác động", "Điều chỉnh mô hình và vận hành"),
], widths=[1.0, 2.05, 2.05, 2.0], font_size=8.4)

doc.add_heading("B. Nguyên tắc dữ liệu có trách nhiệm", level=2)
number(doc, "Thu thập tối thiểu: chỉ lấy dữ liệu cần cho chức năng đã giải thích.")
number(doc, "Đồng thuận rõ ràng: GPS và dữ liệu nhạy cảm phải do người dùng chủ động cấp.")
number(doc, "Phân quyền theo vai trò: dữ liệu cá nhân chi tiết không mặc định công khai.")
number(doc, "Minh bạch mô hình: khuyến nghị phải nêu nguồn dữ liệu, thời điểm và giới hạn.")
number(doc, "Quyền sửa/xóa: người dùng có khả năng xem, cập nhật và yêu cầu xóa dữ liệu theo quy định.")
number(doc, "Không tự động quyết định tín dụng: dữ liệu chỉ hỗ trợ, quyết định thuộc tổ chức được cấp phép và quy trình thẩm định.")

doc.add_heading("C. Nguồn nội bộ dùng để xây dựng hồ sơ", level=2)
bullet(doc, "[STARTUP ZONE 2026] APPLICATION FORM BẢNG A (1).pdf – thông tin đội, dự án, vấn đề, mục tiêu và giai đoạn phát triển.")
bullet(doc, "BÁO CÁO TÀI CHÍNH SUZ2026 (1).xlsx – giả định, ngân lưu, chi phí, thu nhập và chỉ tiêu hiệu quả.")
bullet(doc, "Repo CHAMP MẠNH CHAT tại thời điểm lập hồ sơ – bằng chứng về tính năng MVP và kiến trúc kỹ thuật.")
bullet(doc, "Ảnh hướng dẫn Cách thức dự thi Bảng A – đối chiếu nhóm lĩnh vực và định hướng chủ đề.")
add_p(doc, "Các số liệu bên ngoài được nêu trong hồ sơ này là số liệu đã xuất hiện trong Application Form của đội. Trước khi nộp chính thức, đội cần gắn nguồn gốc/đường dẫn và ngày truy cập theo quy định trích dẫn của Ban Tổ chức.")

doc.add_heading("D. Tuyên bố trung thực về giai đoạn dự án", level=2)
add_p(doc, "CHAMP MẠNH CHAT hiện là MVP đang hoàn thiện và chuẩn bị pilot, chưa thương mại hóa và chưa có đối tác đầu tư chính thức. Hồ sơ phân biệt rõ tính năng đã có, tính năng đang thử nghiệm và năng lực mục tiêu. Các chỉ số tài chính là kịch bản kế hoạch để quản trị và gọi nguồn lực; kết quả thực tế sẽ được cập nhật theo dữ liệu pilot.")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
r = p.add_run("CHAMP MẠNH CHAT")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = RGBColor.from_string(GREEN)
p.add_run("\nĐúng dữ liệu · Đúng quyết định · Đúng nguồn lực")

# Prevent table rows from splitting and normalize table paragraph spacing.
for table in doc.tables:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell in row.cells:
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)

# Document properties
doc.core_properties.title = "CHAMP MẠNH CHAT – File ý tưởng Startup Zone 2026"
doc.core_properties.subject = "Hồ sơ dự thi Bảng A"
doc.core_properties.author = "Đội CHAMP MẠNH CHAT"
doc.core_properties.keywords = "Startup Zone 2026, Bảng A, nông nghiệp, chuyển đổi số, ĐBSCL, xâm nhập mặn"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(str(OUT))
