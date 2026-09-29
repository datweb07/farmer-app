from pathlib import Path
from math import floor

from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"D:\final app\final app")
XLSX = ROOT / "docs" / "BÁO CÁO TÀI CHÍNH FINAL - CHAMP MẠNH CHAT.xlsx"
OUT = ROOT / "docs" / "2.1-5 ADJUST.docx"
TMP = ROOT / ".tmp" / "finance_final" / "adjust_docx"
TMP.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
GREEN = "1F6E52"
LIGHT_BLUE = "EAF2F8"
LIGHT_GREEN = "EAF5EF"
LIGHT_GRAY = "F4F6F7"
MID_GRAY = "D9E1E8"
TEXT = "1F2937"
RED = "C62828"
FONT = "Aptos"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name=FONT, size=10.5, bold=False, color=TEXT, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_pr.append(r_fonts)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_para(doc, text="", bold_lead=None, style=None, align=None, space_after=5, keep=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_font(r, size=10.2)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2, first_col_left=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, text in enumerate(headers):
        cell = hdr.cells[j]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell, 100, 100, 100, 100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(text))
        set_font(r, size=font_size, bold=True, color="FFFFFF")
        if widths:
            cell.width = Cm(widths[j])
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cell = cells[j]
            set_cell_shading(cell, "FFFFFF" if i % 2 == 0 else LIGHT_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell, 90, 100, 90, 100)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 and first_col_left) else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            color = RED if isinstance(value, str) and value.startswith("(") else TEXT
            set_font(r, size=font_size, color=color)
            if widths:
                cell.width = Cm(widths[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, image_path, caption, width_cm=16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(7)
    r = cap.add_run(caption)
    set_font(r, size=9, italic=True, color="555555")


def fmt_money(x):
    if x is None:
        return "-"
    if x < 0:
        return f"({abs(x):,.0f})".replace(",", ".")
    return f"{x:,.0f}".replace(",", ".")


# Read the reviewed model values from the final workbook.
wb = load_workbook(XLSX, data_only=True, read_only=True)
model = wb["Mô hình 5 năm"]
years = [model.cell(4, c).value for c in range(2, 7)]
model_rows = {model.cell(r, 1).value: [model.cell(r, c).value for c in range(2, 7)] for r in range(5, 31)}
cash_ws = wb["Dòng tiền 24 tháng"]
months = [cash_ws.cell(r, 1).value for r in range(5, 29)]
monthly_cash = [cash_ws.cell(r, 16).value for r in range(5, 29)]
fund_ws = wb["Nguồn vốn"]

# Charts from the workbook's reviewed figures, drawn with Pillow so the document
# remains reproducible in the bundled runtime without external chart packages.
FONT_REG = r"C:\Windows\Fonts\arial.ttf"
FONT_BOLD = r"C:\Windows\Fonts\arialbd.ttf"


def fnt(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def money_label(value):
    return f"{value:,.0f}".replace(",", ".")


def make_line_chart(path, title, x_labels, series, y_min, y_max, y_step, highlight_indices=None):
    w, h = 1600, 760
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    left, top, right, bottom = 165, 130, 1510, 630
    d.text((left, 35), title, font=fnt(34, True), fill="#111111")
    d.text((left, 86), "Đơn vị: triệu đồng", font=fnt(20), fill="#555555")
    for y in range(y_min, y_max + 1, y_step):
        py = bottom - (y - y_min) / (y_max - y_min) * (bottom - top)
        d.line((left, py, right, py), fill="#D9E1E8", width=2)
        label = money_label(y)
        box = d.textbbox((0, 0), label, font=fnt(18))
        d.text((left - 20 - (box[2] - box[0]), py - 10), label, font=fnt(18), fill="#555555")
    if y_min < 0 < y_max:
        py0 = bottom - (0 - y_min) / (y_max - y_min) * (bottom - top)
        d.line((left, py0, right, py0), fill="#7F8C8D", width=3)
    n = len(x_labels)
    xs = [left + i * (right - left) / (n - 1) for i in range(n)]
    for i, label in enumerate(x_labels):
        d.text((xs[i] - 24, bottom + 22), str(label), font=fnt(18), fill="#444444")
    legend_x = left
    for name, values, color in series:
        points = []
        for i, val in enumerate(values):
            py = bottom - (val - y_min) / (y_max - y_min) * (bottom - top)
            points.append((xs[i], py))
        d.line(points, fill=color, width=6, joint="curve")
        for i, point in enumerate(points):
            radius = 9 if not highlight_indices or i not in highlight_indices else 13
            d.ellipse((point[0]-radius, point[1]-radius, point[0]+radius, point[1]+radius), fill=color, outline="white", width=3)
        d.line((legend_x, 700, legend_x + 45, 700), fill=color, width=6)
        d.text((legend_x + 58, 685), name, font=fnt(19), fill="#333333")
        legend_x += 310
    im.save(path, quality=95)


chart_5y = TMP / "financial_5y.png"
make_line_chart(
    chart_5y,
    "Doanh thu, chi phí vận hành và EBITDA 5 năm",
    years,
    [
        ("Doanh thu", model_rows["Tổng doanh thu"], "#176B87"),
        ("Chi phí vận hành", model_rows["Tổng chi phí vận hành"], "#5F7F49"),
        ("EBITDA", model_rows["EBITDA"], "#D97925"),
    ],
    -2000, 10000, 2000,
)

chart_cash = TMP / "cash_24m.png"
make_line_chart(
    chart_cash,
    "Số dư tiền mặt 24 tháng đầu",
    [str(i) for i in range(1, 25)],
    [("Tiền cuối kỳ", monthly_cash, "#176B87")],
    0, 600, 100,
    highlight_indices={0, 6, 12},
)

fund_labels = ["Nhà sáng lập", "Grant", "Angel SAFE", "Đối tác chiến lược"]
fund_values = [150, 450, 850, 350]
colors = ["#8AA6B8", "#5FA777", "#176B87", "#D6A84B"]
chart_fund = TMP / "funding_mix.png"
w, h = 1500, 760
im = Image.new("RGB", (w, h), "white")
d = ImageDraw.Draw(im)
d.text((90, 35), "Cơ cấu nguồn vốn 1.800 triệu đồng", font=fnt(34, True), fill="#111111")
box = (110, 150, 680, 720)
start = -90
total = sum(fund_values)
for value, color in zip(fund_values, colors):
    end = start + value / total * 360
    d.pieslice(box, start=start, end=end, fill=color, outline="white", width=4)
    start = end
d.ellipse((260, 300, 530, 570), fill="white")
d.text((335, 388), "1.800", font=fnt(34, True), fill="#17365D")
d.text((340, 435), "triệu", font=fnt(22), fill="#555555")
ly = 205
for label, value, color in zip(fund_labels, fund_values, colors):
    d.rounded_rectangle((790, ly, 835, ly+45), radius=7, fill=color)
    d.text((860, ly-2), f"{label}: {value} triệu ({value/total*100:.1f}%)", font=fnt(23), fill="#222222")
    ly += 105
im.save(chart_fund, quality=95)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)

# Core styles.
styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(TEXT)

title = styles["Title"]
title.font.name = FONT
title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
title.font.size = Pt(22)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = title._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    st = styles[style_name]
    st.font.name = FONT
    st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(9)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True

# Page numbering.
footer = section.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("2.1-5 ADJUST   |   ")
set_font(r, size=8.5, color="666666")
fld_begin = OxmlElement("w:fldChar")
fld_begin.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText")
instr.set(qn("xml:space"), "preserve")
instr.text = "PAGE"
fld_end = OxmlElement("w:fldChar")
fld_end.set(qn("w:fldCharType"), "end")
r._r.append(fld_begin)
r._r.append(instr)
r._r.append(fld_end)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Nội dung điều chỉnh Phần 2.1 và Phần 5")
ppr = p._p.get_or_add_pPr()
p_border = ppr.find(qn("w:pBdr"))
if p_border is not None:
    ppr.remove(p_border)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(14)
r = sub.add_run("Dự án Champ Mạnh Chat   Startup Zone 2026")
set_font(r, size=11, color="555555")
add_para(doc,
         "Tài liệu này gồm hai phần có thể sao chép trực tiếp vào File ý tưởng: phần 2.1 cập nhật vấn đề của nông hộ sầu riêng tại địa bàn thí điểm Cần Thơ và phần 5 trình bày kế hoạch tài chính theo mô hình vốn 1,8 tỷ đồng, giải ngân theo kết quả kiểm chứng.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=10)

doc.add_heading("2.1 Một rủi ro khí hậu kéo theo nhiều quyết định kinh tế", level=1)

add_para(doc,
         "Champ Mạnh Chat lựa chọn thí điểm tại Tây Đô, thành phố Cần Thơ, với trọng tâm là nông hộ trồng sầu riêng. Đây là cây ăn trái có giá trị kinh tế cao nhưng vốn đầu tư lớn, thời gian kiến thiết dài và rất nhạy với chất lượng nguồn nước. Một quyết định lấy nước sai thời điểm có thể làm tăng chi phí phục hồi vườn, ảnh hưởng năng suất và chất lượng trái, sau đó kéo theo thay đổi lịch thu hoạch, kế hoạch thu mua và dòng tiền trả nợ của nông hộ.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_para(doc,
         "Quy mô rủi ro đã đủ lớn để cần một công cụ ra quyết định theo dữ liệu. Theo Báo Nhân Dân ngày 12/5/2026, thành phố Cần Thơ có hơn 14.483 ha sầu riêng; khoảng 8.635 ha đang cho trái và sản lượng năm 2026 ước khoảng 120.000 tấn. Số liệu này phản ánh không gian hành chính Cần Thơ năm 2026. Địa bàn pilot của dự án tập trung vào khu vực Tây Đô và các vùng trồng thuộc Phong Điền, Ô Môn, Thới Lai trước khi mở rộng.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_para(doc,
         "Thiệt hại kinh tế hiện hữu không chỉ đến từ mất mùa. Ngày 21/5/2026, Báo Nhân Dân ghi nhận tại Phong Điền giá thương lái mua tại vườn chỉ còn 20.000-25.000 đồng/kg, so với 50.000-60.000 đồng/kg cùng kỳ năm trước. Chi phí chăm sóc được người trồng phản ánh ở mức khoảng 1,4-1,5 triệu đồng mỗi gốc, khiến nhiều nhà vườn gần như thua lỗ. Với một lô 10 tấn có cùng phẩm cấp, chênh lệch giá nêu trên tương ứng mức hụt doanh thu minh họa khoảng 250-400 triệu đồng trước chi phí. Đây là phép tính minh họa từ mức giá được báo chí ghi nhận, không phải số liệu thiệt hại đã được cơ quan nhà nước kiểm toán.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_para(doc,
         "Rủi ro khí hậu làm biên an toàn của hộ trồng sầu riêng hẹp hơn. Trong cao điểm xâm nhập mặn tháng 3/2026, Báo Cần Thơ dẫn dự báo ranh mặn 4 g/l có thể vào sâu 44-55 km tại các cửa sông Cửu Long. Toàn thành phố có khoảng 8.000 ha cây ăn trái có khả năng chịu ảnh hưởng của hạn hán và xâm nhập mặn. Ngành nông nghiệp khuyến cáo sầu riêng, chôm chôm và măng cụt không được tưới bằng nguồn nước có độ mặn trên 0,5‰. Điều này đặt nông hộ trước một quyết định có tính thời điểm: lấy nước, tiếp tục chờ hay chuyển sang nguồn dự trữ. Nếu thông tin đến chậm hoặc không gắn với vị trí vườn, quyết định có thể sai ngay trong vài giờ quan trọng.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

pain_rows = [
    ("Quy mô ngành hàng tại Cần Thơ năm 2026", "14.483 ha; 8.635 ha cho trái; khoảng 120.000 tấn"),
    ("Giá thu mua tại vườn ở Phong Điền", "20.000-25.000 đồng/kg, so với 50.000-60.000 đồng/kg cùng kỳ trước"),
    ("Chi phí chăm sóc được nhà vườn phản ánh", "Khoảng 1,4-1,5 triệu đồng/gốc"),
    ("Mức hụt doanh thu minh họa cho lô 10 tấn", "Khoảng 250-400 triệu đồng trước chi phí"),
    ("Ngưỡng nước tưới đối với sầu riêng", "Không sử dụng nước có độ mặn trên 0,5‰"),
    ("Diện tích cây ăn trái có nguy cơ hạn mặn toàn thành phố", "Khoảng 8.000 ha trong cao điểm tháng 3/2026"),
]
add_table(doc, ["Chỉ báo pain", "Số liệu cập nhật"], pain_rows, widths=[7.6, 8.4], font_size=9.2)

add_para(doc,
         "Hai nhóm rủi ro phải được quản trị đồng thời nhưng không được đánh đồng. Giá giảm sâu năm 2026 phản ánh khó khăn đầu ra, cung cầu và xuất khẩu; cảnh báo độ mặn phản ánh nguy cơ sản xuất. Khi hai rủi ro xảy ra gần nhau, hộ trồng vừa phải chi thêm để bảo vệ vườn vừa đối diện khả năng bán dưới giá kỳ vọng. Dòng tiền suy giảm làm tăng nguy cơ chậm trả nợ, cắt giảm đầu tư cho vụ sau hoặc bán sớm cho thương lái.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

add_para(doc,
         "Vấn đề cốt lõi vì vậy là khoảng cách giữa dữ liệu và hành động. Thông tin độ mặn, vị trí vườn, nhu cầu thu mua, giá chào và mức độ tin cậy của đối tác đang nằm ở nhiều nguồn. Champ Mạnh Chat chuyển các dữ liệu này thành cảnh báo theo vị trí, khuyến nghị ngắn gọn và kết nối với doanh nghiệp thu mua đã được đánh giá. Kết quả cần đo không phải số lượt xem cảnh báo, mà là tỷ lệ nông hộ ra quyết định đúng thời điểm, số giao dịch hoàn tất, mức sai lệch giữa giá chốt và giá mua thực tế, cùng khả năng duy trì dòng tiền sau mỗi vụ.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

doc.add_heading("Nguồn tham khảo cho Phần 2.1", level=2)
refs = [
    ("[1] Báo Nhân Dân, Cần giải pháp gỡ khó cho ngành hàng sầu riêng, 12/05/2026.", "https://nhandan.vn/can-giai-phap-go-kho-cho-nganh-hang-sau-rieng-post961724.html"),
    ("[2] Báo Nhân Dân, Cần Thơ Nông dân loay hoay tìm đầu ra cho hàng sầu riêng, 21/05/2026.", "https://nhandan.vn/can-tho-nong-dan-loay-hoay-tim-dau-ra-cho-hang-sau-rieng-post963629.html"),
    ("[3] Báo Cần Thơ, Quyết liệt thực hiện các phương án ứng phó hạn xâm nhập mặn, mùa khô 2025-2026.", "https://baocantho.com.vn/quyet-liet-thuc-hien-cac-phuong-an-ung-pho-han-xam-nhap-man-a199606.html"),
    ("[4] Báo Điện tử Chính phủ, Cần Thơ bàn giải pháp nâng cao giá trị xuất khẩu sầu riêng, 21/08/2025.", "https://baochinhphu.vn/can-tho-ban-giai-phap-nang-cao-gia-tri-xuat-khau-sau-rieng-102250821163924929.htm"),
]
for label, url in refs:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    add_hyperlink(p, label, url)

doc.add_page_break()
doc.add_heading("Phần 5 Dự kiến tài chính", level=1)
add_para(doc,
         "Kế hoạch tài chính được xây dựng theo nguyên tắc giải ngân theo kết quả kiểm chứng. Dự án không sử dụng nợ ngân hàng trong 24 tháng đầu. Tổng nhu cầu vốn là 1.800 triệu đồng, thấp hơn 42,9% so với mức 3.154 triệu đồng của mô hình cũ. Vốn được chia thành ba tranche và chỉ giải ngân khi dự án đạt điều kiện về sản phẩm, đối tác pilot, giao dịch hoàn tất và khách hàng trả phí.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

doc.add_heading("5.1 Nguyên tắc và cơ sở lập kế hoạch", level=2)
principles = [
    "Không coi số liệu kế hoạch là doanh thu chắc chắn. Mỗi giả định phải được thay bằng dữ liệu pilot, báo giá, hợp đồng hoặc lịch sử giao dịch khi phát sinh.",
    "Không vay ngân hàng trong 24 tháng đầu vì dự án chưa có doanh thu ổn định, tài sản bảo đảm và khả năng chứng minh DSCR.",
    "Không ghi nhận doanh thu bảo hiểm khi chưa có đối tác được cấp phép và chưa hoàn tất mô hình tuân thủ.",
    "Không mở rộng địa bàn chỉ để tăng người dùng. Dự án ưu tiên mật độ giao dịch, tỷ lệ hoàn tất và khả năng giữ chân doanh nghiệp trả phí.",
    "Nhân sự, hạ tầng và bán hàng chỉ tăng sau khi tranche trước đạt các mốc kiểm chứng đã cam kết.",
]
for item in principles:
    add_bullet(doc, item)

doc.add_heading("5.2 Các giả định tài chính có điều kiện", level=2)
assumption_rows = [
    ("Doanh nghiệp trả phí bình quân", "1,5", "8", "18", "35", "60"),
    ("Phí nền tảng bình quân (triệu đồng/DN/tháng)", "5,0", "6,5", "7,5", "8,0", "8,5"),
    ("GMV giao dịch xác thực (triệu đồng)", "1.800", "9.900", "22.000", "45.000", "69.500"),
    ("Tỷ lệ phí giao dịch", "4,0%", "4,0%", "4,0%", "4,0%", "4,0%"),
    ("Doanh thu pilot onboarding và dữ liệu", "70", "180", "300", "440", "600"),
    ("Thuế thu nhập doanh nghiệp", "20%", "20%", "20%", "20%", "20%"),
    ("Khấu hao", "5 năm", "5 năm", "5 năm", "5 năm", "5 năm"),
    ("Nợ ngân hàng mới", "0", "0", "0", "0", "0"),
]
add_table(doc, ["Giả định", "Năm 1", "Năm 2", "Năm 3", "Năm 4", "Năm 5"], assumption_rows,
          widths=[6.2, 2.0, 2.0, 2.0, 2.0, 2.0], font_size=8.4)
add_para(doc,
         "Mức phí 5 triệu đồng mỗi doanh nghiệp mỗi tháng trong năm đầu là giá thử nghiệm. Mức phí chỉ tăng khi doanh nghiệp xác nhận giá trị tiết kiệm được từ chi phí tìm vùng nguyên liệu, kiểm chứng giao dịch và quản trị lịch thu mua. GMV chỉ bao gồm giao dịch có trạng thái hoàn tất; lead, cuộc gọi và đăng ký chưa phát sinh thu mua không được tính.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

doc.add_page_break()
doc.add_heading("5.3 Kế hoạch doanh thu chi phí và lợi nhuận", level=2)
financial_rows = []
labels = ["Tổng doanh thu", "Tổng chi phí vận hành", "EBITDA", "Biên EBITDA", "Khấu hao", "EBIT", "Thuế TNDN", "Lợi nhuận ròng", "Capex", "Dòng tiền tự do trước tài trợ", "Vốn giải ngân", "Tiền cuối kỳ", "Nợ ngân hàng cuối kỳ"]
for label in labels:
    vals = model_rows[label]
    if label == "Biên EBITDA":
        shown = [f"{v*100:.1f}%" for v in vals]
    else:
        shown = [fmt_money(v) for v in vals]
    financial_rows.append(tuple([label] + shown))
add_table(doc, ["Chỉ tiêu triệu đồng", "Năm 1", "Năm 2", "Năm 3", "Năm 4", "Năm 5"], financial_rows,
          widths=[5.9, 2.05, 2.05, 2.05, 2.05, 2.05], font_size=8.3)
add_figure(doc, chart_5y, "Hình 1. Doanh thu, chi phí vận hành và EBITDA theo mô hình tài chính điều chỉnh", width_cm=13.6)

add_para(doc,
         "Dự án chấp nhận EBITDA âm 887 triệu đồng trong Năm 1 và âm 300 triệu đồng trong Năm 2 để hoàn thiện MVP, vận hành pilot và xây dựng mật độ giao dịch. EBITDA chuyển dương 550 triệu đồng trong Năm 3. Lợi nhuận ròng chuyển dương cùng năm, đạt khoảng 333 triệu đồng. Đến Năm 5, doanh thu đạt 9.500 triệu đồng, EBITDA đạt 4.250 triệu đồng và lợi nhuận ròng đạt khoảng 3.197 triệu đồng. Đây là kịch bản cơ sở có điều kiện, không phải cam kết doanh thu.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

doc.add_heading("5.4 Cấu trúc doanh thu", level=2)
revenue_rows = [
    ("Thuê bao nền tảng", *[fmt_money(x) for x in model_rows["Doanh thu thuê bao"]]),
    ("Phí giao dịch", *[fmt_money(x) for x in model_rows["Doanh thu phí giao dịch"]]),
    ("Pilot onboarding và dữ liệu", *[fmt_money(x) for x in model_rows["Doanh thu pilot, onboarding và dữ liệu"]]),
    ("Tổng doanh thu", *[fmt_money(x) for x in model_rows["Tổng doanh thu"]]),
]
add_table(doc, ["Nguồn doanh thu triệu đồng", "Năm 1", "Năm 2", "Năm 3", "Năm 4", "Năm 5"], revenue_rows,
          widths=[5.9, 2.05, 2.05, 2.05, 2.05, 2.05], font_size=8.5)
add_para(doc,
         "Thuê bao nền tảng tạo phần doanh thu định kỳ; phí giao dịch tăng theo GMV đã xác thực; doanh thu pilot, onboarding và dữ liệu phản ánh dịch vụ triển khai cho đối tác. Mô hình không ghi nhận doanh thu từ việc bán dữ liệu cá nhân và không đưa doanh thu bảo hiểm vào kịch bản cơ sở.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

doc.add_heading("5.5 Cấu trúc chi phí vận hành", level=2)
cost_rows = [
    ("Nhân sự", *[fmt_money(x) for x in model_rows["Nhân sự"]]),
    ("Cloud dữ liệu và bảo mật", *[fmt_money(x) for x in model_rows["Cloud, dữ liệu và bảo mật"]]),
    ("Pilot và vận hành địa bàn", *[fmt_money(x) for x in model_rows["Pilot và vận hành địa bàn"]]),
    ("Bán hàng và marketing", *[fmt_money(x) for x in model_rows["Bán hàng và marketing"]]),
    ("Pháp lý và hành chính", *[fmt_money(x) for x in model_rows["Pháp lý và hành chính"]]),
    ("Tổng chi phí vận hành", *[fmt_money(x) for x in model_rows["Tổng chi phí vận hành"]]),
]
add_table(doc, ["Khoản mục triệu đồng", "Năm 1", "Năm 2", "Năm 3", "Năm 4", "Năm 5"], cost_rows,
          widths=[5.9, 2.05, 2.05, 2.05, 2.05, 2.05], font_size=8.5)
add_para(doc,
         "Nhân sự là khoản chi lớn nhất vì dự án cần đồng thời phát triển sản phẩm, vận hành dữ liệu và hỗ trợ pilot tại địa bàn. Chi phí cloud và bảo mật tăng theo lượng giao dịch và yêu cầu lưu vết. Chi phí bán hàng chỉ mở rộng sau khi nhóm chứng minh được khả năng chuyển đổi doanh nghiệp pilot thành khách hàng trả phí.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

doc.add_heading("5.6 Phương án huy động vốn theo tranche", level=2)
tranche_rows = [
    ("Tranche 0", "Tháng 1", "700", "150 nhà sáng lập; 250 grant; 300 angel", "MVP ổn định; 3 đối tác pilot; danh sách 100 nông hộ tiềm năng"),
    ("Tranche 1", "Tháng 7", "600", "150 grant; 300 angel; 150 đối tác", "Tối thiểu 10 giao dịch hoàn tất; 2 doanh nghiệp bắt đầu trả phí; có dữ liệu pilot"),
    ("Tranche 2", "Tháng 13", "500", "50 grant; 250 angel; 200 đối tác", "4 doanh nghiệp trả phí; completion từ 70%; retention pilot từ 60%"),
]
add_table(doc, ["Tranche", "Thời điểm", "Số tiền triệu", "Nguồn", "Điều kiện giải ngân"], tranche_rows,
          widths=[2.3, 2.2, 2.1, 4.3, 6.3], font_size=8.5)

source_rows = [
    ("Nhà sáng lập", "150", "8,3%"),
    ("Grant từ cuộc thi vườn ươm chương trình đổi mới", "450", "25,0%"),
    ("Angel SAFE", "850", "47,2%"),
    ("Đối tác chiến lược", "350", "19,4%"),
    ("Tổng", "1.800", "100,0%"),
]
add_table(doc, ["Nguồn vốn", "Số tiền triệu đồng", "Tỷ trọng"], source_rows, widths=[9.2, 3.6, 3.2], font_size=9)
add_figure(doc, chart_fund, "Hình 2. Cơ cấu huy động vốn trong 24 tháng đầu", width_cm=14.7)

doc.add_page_break()
doc.add_heading("Phân bổ sử dụng vốn", level=3)
use_rows = [
    ("Sản phẩm dữ liệu và Capex", "490", "27,2%", "Bảo mật, dữ liệu, hạ tầng và sản phẩm sau MVP"),
    ("Cầu nối chi phí nhân sự", "620", "34,4%", "Giữ đội ngũ tinh gọn trong 24 tháng đầu"),
    ("Pilot và vận hành địa bàn", "320", "17,8%", "Onboarding, hỗ trợ, đo lường kết quả mùa vụ"),
    ("Bán hàng và phát triển đối tác", "160", "8,9%", "Pipeline B2B, hợp đồng pilot và tài liệu bán hàng"),
    ("Pháp lý và hành chính", "120", "6,7%", "Dữ liệu cá nhân, hợp đồng, kế toán và tuân thủ"),
    ("Dự phòng", "90", "5,0%", "Biến động dữ liệu, tích hợp và thời vụ"),
    ("Tổng", "1.800", "100,0%", ""),
]
add_table(doc, ["Hạng mục sử dụng vốn", "Số tiền triệu", "Tỷ trọng", "Mục đích"], use_rows,
          widths=[5.2, 2.7, 2.2, 6.1], font_size=8.6)

doc.add_heading("5.7 Dòng tiền 24 tháng và runway", level=2)
add_figure(doc, chart_cash, "Hình 3. Số dư tiền mặt theo tháng sau ba đợt giải ngân", width_cm=16.1)
cash_rows = [
    ("Tổng doanh thu 24 tháng", "1.432 triệu đồng"),
    ("Tổng chi phí vận hành 24 tháng", "2.619 triệu đồng"),
    ("Capex 24 tháng", "490 triệu đồng"),
    ("Tổng vốn giải ngân", "1.800 triệu đồng"),
    ("Số dư thấp nhất trước tranche tiếp theo", "41 triệu đồng cuối Tháng 6"),
    ("Tiền cuối Năm 1", "73 triệu đồng"),
    ("Tiền cuối Năm 2", "123 triệu đồng"),
    ("Nợ ngân hàng", "0 đồng"),
]
add_table(doc, ["Chỉ tiêu runway", "Kết quả"], cash_rows, widths=[9.2, 6.8], font_size=9.2)
add_para(doc,
         "Tranche 1 phải hoàn tất cam kết trước cuối Tháng 6 để tiền sẵn sàng vào đầu Tháng 7. Mức đệm 41 triệu đồng là thấp, vì vậy nhóm phải theo dõi dòng tiền hàng tuần và dừng tuyển dụng hoặc capex chưa thiết yếu nếu việc giải ngân chậm. Kế hoạch chỉ xem xét hạn mức tín dụng sau 24 tháng, khi MRR, hợp đồng hoặc khoản phải thu và khả năng trả nợ đã có thể kiểm chứng.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

doc.add_heading("5.8 Phân tích độ nhạy", level=2)
sensitivity_rows = [
    ("Downside", "70%", "95%", "13.532", "642", "(176)", "176", "Năm 4"),
    ("Cơ sở", "100%", "100%", "19.332", "5.763", "123", "0", "Năm 3"),
    ("Upside", "130%", "110%", "25.132", "10.206", "291", "0", "Năm 3"),
]
add_table(doc,
          ["Kịch bản", "Hệ số doanh thu", "Hệ số Opex", "Doanh thu 5 năm", "EBITDA 5 năm", "Tiền cuối Năm 2", "Vốn bổ sung", "EBITDA dương"],
          sensitivity_rows, widths=[2.3, 2.0, 1.8, 2.5, 2.3, 2.3, 2.0, 2.0], font_size=7.7)
add_para(doc,
         "Kịch bản downside là ngưỡng quản trị, không phải lời tiên đoán. Nếu doanh thu chỉ đạt 70% kế hoạch trong khi chi phí chỉ giảm được 5%, dự án thiếu khoảng 176 triệu đồng vào cuối Năm 2 và EBITDA dương bị lùi sang Năm 4. Khi doanh thu thực tế thấp hơn kế hoạch trong hai quý liên tiếp, nhóm phải giảm tốc mở rộng, hoãn capex và chuẩn bị bridge round trước khi runway xuống dưới ba tháng.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

doc.add_heading("5.9 Cổng kiểm chứng và kỷ luật tài chính", level=2)
gate_rows = [
    ("Sản phẩm và pilot", "MVP ổn định; 3 đối tác pilot; 100 nông hộ trong pipeline", "Mở Tranche 0"),
    ("Giao dịch", "Tối thiểu 10 giao dịch hoàn tất; dữ liệu giá và trạng thái được lưu", "Cho phép tiếp tục pilot trả phí"),
    ("Khách hàng trả phí", "2 doanh nghiệp bắt đầu trả phí trước Tranche 1; 4 doanh nghiệp trước Tranche 2", "Cho phép tăng nhân sự và bán hàng"),
    ("Chất lượng giao dịch", "Tỷ lệ hoàn tất từ 70%; retention pilot từ 60%", "Cho phép mở rộng địa bàn"),
    ("Runway", "Cập nhật 13 tuần; cảnh báo khi còn dưới 3 tháng tiền mặt", "Dừng capex không thiết yếu và chuẩn bị bridge"),
    ("Vốn vay", "Chỉ xem xét sau 24 tháng khi có MRR ổn định, hợp đồng hoặc khoản phải thu và DSCR kiểm chứng được", "Không dùng nợ để bù mô hình chưa chứng minh"),
]
add_table(doc, ["Cổng", "Bằng chứng yêu cầu", "Quyết định tài chính"], gate_rows,
          widths=[3.6, 8.0, 4.6], font_size=8.5)

add_para(doc,
         "Trong giai đoạn pre-revenue, dự án không sử dụng NPV và IRR làm bằng chứng chính cho tính khả thi. Hai chỉ tiêu này chỉ được tính lại sau 12-18 tháng khi đã có dữ liệu về tỷ lệ chuyển đổi trả phí, doanh thu giữ lại, chi phí phục vụ một doanh nghiệp và dòng tiền giao dịch thực tế. Trọng tâm của kế hoạch hiện tại là duy trì runway, kiểm chứng willingness-to-pay và chỉ giải ngân khi đạt mốc vận hành.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

doc.add_heading("5.10 Kết luận tài chính", level=2)
add_para(doc,
         "Phương án điều chỉnh giảm nhu cầu vốn xuống 1.800 triệu đồng và loại bỏ nợ ngân hàng trong 24 tháng đầu. Cấu trúc này phù hợp hơn với một startup công nghệ nông nghiệp đang kiểm chứng thị trường: rủi ro của nhà sáng lập và nhà đầu tư được giới hạn theo từng tranche; chi phí chỉ tăng khi có bằng chứng sử dụng; và dự án vẫn giữ được khả năng hòa vốn EBITDA trong Năm 3 ở kịch bản cơ sở. Nếu các chỉ số trả phí, hoàn tất giao dịch hoặc giữ chân không đạt mốc, dự án phải thu hẹp phạm vi và điều chỉnh chi phí trước khi huy động vòng tiếp theo.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=10)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
r = p.add_run("Nguồn số liệu Phần 5: BÁO CÁO TÀI CHÍNH FINAL - CHAMP MẠNH CHAT.xlsx, phiên bản đã đối chiếu ngày 25/09/2026. Đơn vị tiền tệ trong các bảng là triệu đồng, trừ khi có ghi chú khác.")
set_font(r, size=9, italic=True, color="555555")

# Prevent table rows from splitting where possible and apply paragraph orphan control.
for table in doc.tables:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_together = True

for p in doc.paragraphs:
    p_pr = p._p.get_or_add_pPr()
    widow = OxmlElement("w:widowControl")
    widow.set(qn("w:val"), "1")
    p_pr.append(widow)

doc.core_properties.title = "Nội dung điều chỉnh Phần 2.1 và Phần 5"
doc.core_properties.subject = "Dự án Champ Mạnh Chat Startup Zone 2026"
doc.core_properties.author = "Nhóm dự án Champ Mạnh Chat"
doc.save(OUT)
print(OUT)
